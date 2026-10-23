"""Coordinate-based Office operations, using mature parsers and atomic saves."""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path

from agent_runtime.workspace import atomic_bytes, check_version, file_version


def _value(value):
    return value.isoformat() if isinstance(value, (date, datetime, time)) else value


@lru_cache(maxsize=24)
def _scan(path: str, version: str):
    from tools.workbook_scan import scan_workbook

    return scan_workbook(path)


def inspect_workbook(path: Path):
    version = file_version(path)
    scan = _scan(str(path), version)
    return {"success": True, "version": version, "workbook": scan}


def read_sheet_range(path: Path, sheet: str, cell_range: str, formulas: bool):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries, get_column_letter

    version = file_version(path)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    if (
        not all((min_col, min_row, max_col, max_row))
        or min_col > max_col
        or min_row > max_row
    ):
        raise ValueError("范围必须使用完整 A1 地址，例如 A1:D30")
    if (max_col - min_col + 1) * (max_row - min_row + 1) > 2000:
        raise ValueError("一次最多读取 2000 个单元格，请缩小区域")
    cells = []
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        try:
            ws = book.sheet_by_name(sheet)
            for row in range(min_row - 1, min(max_row, ws.nrows)):
                for col in range(min_col - 1, min(max_col, ws.ncols)):
                    cell = ws.cell(row, col)
                    value = (
                        xlrd.xldate_as_datetime(cell.value, book.datemode)
                        if cell.ctype == xlrd.XL_CELL_DATE
                        else cell.value
                    )
                    cells.append(
                        {
                            "coordinate": f"{get_column_letter(col + 1)}{row + 1}",
                            "value": _value(value),
                            "formula": None,
                            "number_format": None,
                        }
                    )
        finally:
            book.release_resources()
        warnings = ["旧版 XLS 仅提供缓存值，不提供公式正文"]
    else:
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        source = None
        try:
            if sheet not in values.sheetnames:
                raise ValueError("工作表不存在")
            source = (
                openpyxl.load_workbook(path, read_only=True, data_only=False)
                if formulas
                else None
            )
            from itertools import repeat

            bounds = dict(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
            )
            formula_rows = source[sheet].iter_rows(**bounds) if source else repeat(None)
            for row_index, (row, formula_row) in enumerate(
                zip(values[sheet].iter_rows(**bounds), formula_rows), min_row
            ):
                for col_index, cell in enumerate(row, min_col):
                    coordinate = f"{get_column_letter(col_index)}{row_index}"
                    original = formula_row[col_index - min_col] if formula_row else None
                    formula = (
                        original.value
                        if original and original.data_type == "f"
                        else None
                    )
                    cells.append(
                        {
                            "coordinate": coordinate,
                            "value": _value(cell.value),
                            "formula": formula,
                            "number_format": getattr(cell, "number_format", "General"),
                        }
                    )
            warnings = [
                "公式结果为工作簿缓存值；缺失或过期缓存需要通过导入复核/重算处理"
            ]
        finally:
            values.close()
            if source:
                source.close()
    check_version(path, version)
    return {
        "success": True,
        "sheet": sheet,
        "range": cell_range,
        "cells": cells,
        "version": version,
        "warnings": warnings,
    }


def edit_workbook(path: Path, expected_version: str, changes: list[dict]):
    import openpyxl
    from openpyxl.utils.cell import coordinate_to_tuple

    check_version(path, expected_version)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("精确编辑支持 XLSX/XLSM；旧 XLS 请先通过导出服务转换副本")
    if path.exists():
        with zipfile.ZipFile(path) as package:
            unsupported = [
                name
                for name in package.namelist()
                if name.startswith(
                    ("xl/drawings/", "xl/activeX/", "xl/embeddings/", "xl/pivot")
                )
            ]
            if unsupported:
                raise ValueError(
                    "文件含图形/嵌入对象/透视对象，当前编辑器无法保证保留；请用转换副本或原生 Office 处理"
                )
        book = openpyxl.load_workbook(
            path,
            keep_vba=path.suffix.lower() == ".xlsm",
            keep_links=True,
            rich_text=True,
        )
    else:
        book = openpyxl.Workbook()
    try:
        for change in changes:
            row, col = coordinate_to_tuple(change["cell"])
            if row > 1048576 or col > 16384:
                raise ValueError("单元格地址超出 Excel 范围")
            sheet = change["sheet"]
            ws = book[sheet] if sheet in book.sheetnames else book.create_sheet(sheet)
            value = change["value"]
            if change["kind"] == "formula":
                if not isinstance(value, str) or not value.startswith("="):
                    raise ValueError("公式必须以 = 开头")
            elif isinstance(value, str) and value.startswith("="):
                ws.cell(row, col, value)
                ws.cell(row, col).data_type = "s"
                continue
            ws.cell(row, col).value = value
        if (
            changes
            and expected_version == "missing"
            and "Sheet" in book.sheetnames
            and len(book.sheetnames) > 1
        ):
            if book["Sheet"].max_row == 1 and book["Sheet"]["A1"].value is None:
                del book["Sheet"]
        buffer = io.BytesIO()
        book.save(buffer)
    finally:
        book.close()
    # Reopen the produced archive before replacing the user's version.
    verified = openpyxl.load_workbook(
        io.BytesIO(buffer.getvalue()),
        read_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )
    verified.close()
    version = atomic_bytes(path, buffer.getvalue(), expected_version)
    return {"success": True, "version": version, "changed_cells": len(changes)}


def _blocks(document):
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    def walk(parent, prefix):
        for index, item in enumerate(parent.iter_inner_content()):
            identifier = f"{prefix}/{index}"
            if isinstance(item, Paragraph):
                yield identifier, item
            elif isinstance(item, Table):
                seen = set()
                for row_index, row in enumerate(item.rows):
                    for col_index, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        yield from walk(cell, f"{identifier}/r{row_index}/c{col_index}")

    yield from walk(document, "body")
    for index, section in enumerate(document.sections):
        yield from walk(section.header, f"header{index}")
        yield from walk(section.footer, f"footer{index}")


def read_document(path: Path, offset: int = 0, limit: int = 200, character_offset: int = 0):
    if offset < 0 or limit < 1 or character_offset < 0:
        raise ValueError("分页位置或数量无效")
    version = file_version(path)
    if path.suffix.lower() in {".md", ".txt", ".csv", ".json"}:
        with path.open(encoding="utf-8") as stream:
            remaining = character_offset
            while remaining:
                skipped = stream.read(min(remaining, 65536))
                if not skipped:
                    raise ValueError("字符续读位置超出文档范围")
                remaining -= len(skipped)
            content = stream.read(20001)
        check_version(path, version)
        return {
            "success": True,
            "version": version,
            "text": content[:20000],
            "truncated": len(content) > 20000,
            "next_offset": 0,
            "next_character_offset": character_offset + min(len(content), 20000),
        }
    if path.suffix.lower() != ".docx":
        raise ValueError("结构化读取支持 DOCX 和文本；其他格式请先转换为副本")
    from docx import Document

    document = Document(path)
    blocks = list(_blocks(document))
    selected = []
    size = 0
    next_offset, next_character_offset = offset, 0
    for index, (identifier, paragraph) in enumerate(blocks[offset : offset + limit], offset):
        start = character_offset if index == offset else 0
        if start > len(paragraph.text):
            raise ValueError("字符续读位置超出当前文档块范围")
        text = paragraph.text[start:]
        if size + len(text) > 20000:
            if not selected:
                selected.append(
                    {
                        "id": identifier,
                        "text": text[:20000],
                        "truncated": True,
                        "style": paragraph.style.name if paragraph.style else "",
                    }
                )
                next_character_offset = start + 20000
            break
        selected.append(
            {
                "id": identifier,
                "text": text,
                "style": paragraph.style.name if paragraph.style else "",
            }
        )
        size += len(text)
        next_offset = index + 1
    check_version(path, version)
    return {
        "success": True,
        "version": version,
        "blocks": selected,
        "next_offset": next_offset,
        "next_character_offset": next_character_offset,
        "truncated": next_offset < len(blocks),
        "warnings": ["返回正文、表格、页眉页脚的文本；图片、复杂域及修订需专门处理"],
    }


def edit_document(path: Path, expected_version: str, changes: list[dict]):
    from docx import Document

    check_version(path, expected_version)
    if path.suffix.lower() != ".docx":
        raise ValueError("结构化文档编辑只支持 DOCX")
    document = Document(path)
    paragraphs = dict(_blocks(document))
    for change in changes:
        paragraph = paragraphs.get(change["block_id"])
        if paragraph is None:
            raise ValueError("文档块不存在，请重新读取文档")
        if any(
            paragraph._p.xpath(f".//w:{tag}")
            for tag in ("fldChar", "ins", "del", "hyperlink")
        ):
            raise ValueError("目标段落含复杂域/修订/链接，不能进行保真文本替换")
        old, new = change["old_text"], change["new_text"]
        text = "".join(run.text for run in paragraph.runs)
        if not old or text.count(old) != 1:
            raise ValueError("原文必须在指定块中唯一匹配")
        start = text.index(old)
        end = start + len(old)
        position = 0
        inserted = False
        for run in paragraph.runs:
            current = run.text
            stop = position + len(current)
            if position < end and stop > start:
                if run._r.xpath(".//w:drawing"):
                    raise ValueError("目标文字与图片处于同一 run，无法保真修改")
                prefix = current[: max(0, start - position)]
                suffix = current[max(0, end - position) :] if end < stop else ""
                run.text = prefix + (new if not inserted else "") + suffix
                inserted = True
            position = stop
    buffer = io.BytesIO()
    document.save(buffer)
    Document(io.BytesIO(buffer.getvalue()))
    version = atomic_bytes(path, buffer.getvalue(), expected_version)
    return {"success": True, "version": version, "changed_blocks": len(changes)}
