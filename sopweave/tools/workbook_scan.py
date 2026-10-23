"""Workbook structure and formula-cache scanner.

This module is the read-only front door for Excel ingestion. It records what
is physically present before layout inference or materialization discards it.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from config.brand import CACHE_DIR_NAME, LEGACY_CACHE_DIR_NAME
from config.settings import IMPORT_MAX_CELLS

MANIFEST_SCHEMA_VERSION = 3
IMPORT_CACHE_DIR_NAME = "import-cache"
FORMULA_EXAMPLE_LIMIT = 20
MAX_IMPORT_CELLS = IMPORT_MAX_CELLS
ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
}


class WorkbookScanError(ValueError):
    """Raised when a workbook cannot be scanned safely."""


def read_workbook_values(
    path: str | os.PathLike[str],
    max_rows: int | None = None,
    *,
    manifest: dict[str, Any] | None = None,
) -> dict[str, list[list[Any]]]:
    """Read cached/displayed values through one shared workbook adapter."""
    source = Path(path)
    metadata = manifest if manifest is not None else scan_workbook(source)
    bounds = {
        sheet["name"]: sheet.get("actual_bounds")
        for sheet in metadata.get("sheets", [])
    }
    cells = sum(region[2] * region[3] for region in bounds.values() if region)
    if MAX_IMPORT_CELLS and cells > MAX_IMPORT_CELLS:
        raise WorkbookScanError(
            f"工作簿有效区域 {cells} 个单元格超过 {MAX_IMPORT_CELLS} 单元格预算，未截断导入"
        )
    if max_rows is not None and any(
        region and region[2] > max_rows for region in bounds.values()
    ):
        raise WorkbookScanError(f"工作簿超过单表 {max_rows} 行上限，未截断导入")
    if source.suffix.lower() == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(source)
        result: dict[str, list[list[Any]]] = {}
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            region = bounds.get(sheet_name)
            stop = region[2] if region else 0
            rows: list[list[Any]] = []
            for row_index in range(stop):
                values: list[Any] = []
                for column_index in range(region[3]):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        except (TypeError, ValueError):
                            pass
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif isinstance(value, float) and value.is_integer():
                        value = int(value)
                    values.append(value)
                rows.append(values)
            result[sheet_name] = rows
        return result

    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    result = {}
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            region = bounds.get(sheet_name)
            if not region:
                result[sheet_name] = []
                continue
            rows = []
            for raw in worksheet.iter_rows(
                min_row=1,
                min_col=1,
                max_row=region[2],
                max_col=region[3],
                values_only=True,
            ):
                rows.append(list(raw))
            result[sheet_name] = rows
    finally:
        workbook.close()
    return result


def sha256_path(path: str | os.PathLike[str]) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_cache_root(workspace: str | os.PathLike[str]) -> Path:
    """Return the current cache root and migrate the pre-rename directory once."""
    root = Path(workspace)
    current = root / CACHE_DIR_NAME
    legacy = root / LEGACY_CACHE_DIR_NAME
    if legacy.is_dir() and not current.exists():
        legacy.rename(current)
    return current


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return value


def _cell_range(bounds: tuple[int, int, int, int] | None) -> str:
    if bounds is None:
        return ""
    from openpyxl.utils import get_column_letter

    min_row, min_col, max_row, max_col = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _actual_bounds(cells: Iterable[Any]) -> tuple[int, int, int, int] | None:
    meaningful = [
        cell
        for cell in cells
        if getattr(cell, "value", None) is not None
        and getattr(cell, "value", None) != ""
    ]
    if not meaningful:
        return None
    return (
        min(cell.row for cell in meaningful),
        min(cell.column for cell in meaningful),
        max(cell.row for cell in meaningful),
        max(cell.column for cell in meaningful),
    )


def _stored_cells(worksheet: Any) -> list[Any]:
    """Return physically stored cells without expanding a stale dimension."""
    sparse = getattr(worksheet, "_cells", None)
    if isinstance(sparse, dict):
        return list(sparse.values())
    return [
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None and cell.value != ""
    ]


def _is_error(value: Any, data_type: str | None = None) -> bool:
    return data_type == "e" or (
        isinstance(value, str) and value.upper() in ERROR_VALUES
    )


def _is_direct_self_reference(formula: str, coordinate: str, sheet_name: str) -> bool:
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import TokenizerError
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

    row, column = coordinate_to_tuple(coordinate.replace("$", ""))
    try:
        tokens = Tokenizer(formula).items
    except TokenizerError:
        return False
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        address = token.value
        if "!" in address:
            qualifier, address = address.rsplit("!", 1)
            qualifier = qualifier.strip("'").replace("''", "'")
            if qualifier.casefold() != sheet_name.casefold():
                continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except ValueError:
            continue
        if (min_col or 1) <= column <= (max_col or 16384) and (min_row or 1) <= row <= (
            max_row or 1048576
        ):
            return True
    return False


def _calculation_properties(workbook: Any) -> dict[str, Any]:
    calculation = getattr(workbook, "calculation", None)
    if calculation is None:
        return {}
    fields = (
        "calcId",
        "calcMode",
        "fullCalcOnLoad",
        "forceFullCalc",
        "iterate",
        "iterateCount",
        "iterateDelta",
    )
    return {
        name: _json_value(getattr(calculation, name))
        for name in fields
        if getattr(calculation, name, None) is not None
    }


def _defined_names(workbook: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    values = getattr(getattr(workbook, "defined_names", None), "values", None)
    if values is None:
        return result
    for item in values():
        name = str(getattr(item, "name", "") or "")
        target = str(getattr(item, "attr_text", "") or "")
        if name:
            destinations: list[dict[str, str]] = []
            try:
                destinations = [
                    {"sheet": str(sheet), "ref": str(ref)}
                    for sheet, ref in item.destinations
                ]
            except (AttributeError, TypeError, ValueError):
                pass
            result.append(
                {
                    "name": name,
                    "target": target,
                    "destinations": destinations,
                }
            )
    return sorted(result, key=lambda item: item["name"].casefold())


def _scan_ooxml(path: Path, file_hash: str) -> dict[str, Any]:
    import openpyxl

    with path.open("rb") as stream:
        compound_header = stream.read(8) == bytes.fromhex("d0cf11e0a1b11ae1")
    if compound_header:
        import io
        from xlrd.compdoc import CompDoc, CompDocError

        try:
            document = CompDoc(path.read_bytes(), logfile=io.StringIO())
            encrypted = document.get_named_stream("EncryptionInfo") is not None
        except CompDocError:
            encrypted = False
        if encrypted:
            raise WorkbookScanError(
                "工作簿已加密，当前不支持直接导入；请先在 Excel 中解密并另存为 .xlsx 后重试"
            )

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    try:
        formula_wb = openpyxl.load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )
        value_wb = openpyxl.load_workbook(
            path,
            read_only=False,
            data_only=True,
            keep_vba=keep_vba,
        )
    except Exception as exc:  # noqa: BLE001
        raise WorkbookScanError(f"无法读取 Excel 工作簿: {exc}") from exc

    calculation_properties = _calculation_properties(formula_wb)
    defined_names = _defined_names(formula_wb)
    sheets: list[dict[str, Any]] = []
    try:
        for index, formula_ws in enumerate(formula_wb.worksheets):
            value_ws = value_wb[formula_ws.title]
            cells = _stored_cells(formula_ws)
            bounds = _actual_bounds(cells)
            formulas = [
                cell for cell in cells if getattr(cell, "data_type", None) == "f"
            ]
            examples: list[dict[str, Any]] = []
            cached_count = missing_count = formula_error_count = 0
            cross_sheet_count = external_count = direct_self_count = 0
            for cell in formulas:
                formula = str(cell.value or "")
                cached = value_ws[cell.coordinate].value
                missing = cached is None or cached == ""
                cached_count += 0 if missing else 1
                missing_count += 1 if missing else 0
                formula_error_count += 1 if _is_error(cached) else 0
                cross_sheet = "!" in formula
                external = "[" in formula and "]" in formula
                direct_self = _is_direct_self_reference(
                    formula, cell.coordinate, formula_ws.title
                )
                cross_sheet_count += 1 if cross_sheet else 0
                external_count += 1 if external else 0
                direct_self_count += 1 if direct_self else 0
                if len(examples) < FORMULA_EXAMPLE_LIMIT:
                    examples.append(
                        {
                            "cell": cell.coordinate,
                            "formula": formula,
                            "cached_value": _json_value(cached),
                            "missing_cached_value": missing,
                            "cross_sheet_reference": cross_sheet,
                            "external_reference": external,
                            "direct_self_reference": direct_self,
                        }
                    )

            error_cells = []
            for cell in cells:
                cached = value_ws[cell.coordinate].value
                if _is_error(cell.value, getattr(cell, "data_type", None)) or _is_error(
                    cached
                ):
                    error_cells.append(cell.coordinate)

            tables = sorted(
                (
                    {"name": str(table.name), "ref": str(table.ref)}
                    for table in formula_ws.tables.values()
                ),
                key=lambda item: item["name"].casefold(),
            )
            sheets.append(
                {
                    "name": formula_ws.title,
                    "index": index,
                    "state": formula_ws.sheet_state,
                    "declared_range": formula_ws.calculate_dimension(),
                    "actual_range": _cell_range(bounds),
                    "actual_bounds": list(bounds) if bounds else None,
                    "stored_cell_count": len(cells),
                    "non_empty_cell_count": sum(
                        1
                        for cell in cells
                        if cell.value is not None and cell.value != ""
                    ),
                    "formula_count": len(formulas),
                    "formula_cached_count": cached_count,
                    "formula_missing_count": missing_count,
                    "formula_error_count": formula_error_count,
                    "cross_sheet_formula_count": cross_sheet_count,
                    "external_formula_count": external_count,
                    "direct_self_reference_count": direct_self_count,
                    "formula_examples": examples,
                    "error_cell_count": len(error_cells),
                    "error_cells": error_cells[:FORMULA_EXAMPLE_LIMIT],
                    "hidden_rows": sorted(
                        int(row_index)
                        for row_index, dimension in formula_ws.row_dimensions.items()
                        if dimension.hidden
                    ),
                    "hidden_columns": sorted(
                        str(column)
                        for column, dimension in formula_ws.column_dimensions.items()
                        if dimension.hidden
                    ),
                    "merged_ranges": sorted(
                        str(value) for value in formula_ws.merged_cells.ranges
                    ),
                    "tables": tables,
                }
            )
    finally:
        formula_wb.close()
        value_wb.close()

    formula_count = sum(sheet["formula_count"] for sheet in sheets)
    formula_missing = sum(sheet["formula_missing_count"] for sheet in sheets)
    return {
        "schema_version": MANIFEST_SCHEMA_VERSION,
        "file_hash": file_hash,
        "source_name": path.name,
        "format": path.suffix.lower().lstrip("."),
        "file_size": path.stat().st_size,
        "calculation_properties": calculation_properties,
        "defined_names": defined_names,
        "sheet_count": len(sheets),
        "formula_count": formula_count,
        "formula_cached_count": sum(sheet["formula_cached_count"] for sheet in sheets),
        "formula_missing_count": formula_missing,
        "formula_error_count": sum(sheet["formula_error_count"] for sheet in sheets),
        "has_unresolved_formulas": formula_missing > 0,
        "sheets": sheets,
    }


def _scan_xls(path: Path, file_hash: str) -> dict[str, Any]:
    try:
        import xlrd

        workbook = xlrd.open_workbook(path)
    except Exception as exc:  # noqa: BLE001
        raise WorkbookScanError(f"无法读取旧版 Excel 工作簿: {exc}") from exc

    sheets: list[dict[str, Any]] = []
    for index, worksheet in enumerate(workbook.sheets()):
        min_row = min_col = None
        max_row = max_col = 0
        error_cells: list[str] = []
        non_empty = 0
        for row_index in range(worksheet.nrows):
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                non_empty += 1
                row_number = row_index + 1
                column_number = column_index + 1
                min_row = row_number if min_row is None else min(min_row, row_number)
                min_col = (
                    column_number if min_col is None else min(min_col, column_number)
                )
                max_row = max(max_row, row_number)
                max_col = max(max_col, column_number)
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    from openpyxl.utils import get_column_letter

                    error_cells.append(
                        f"{get_column_letter(column_number)}{row_number}"
                    )
        bounds = None if min_row is None else (min_row, min_col, max_row, max_col)
        sheets.append(
            {
                "name": worksheet.name,
                "index": index,
                "state": (
                    "visible" if getattr(worksheet, "visibility", 0) == 0 else "hidden"
                ),
                "declared_range": _cell_range(bounds),
                "actual_range": _cell_range(bounds),
                "actual_bounds": list(bounds) if bounds else None,
                "stored_cell_count": worksheet.nrows * worksheet.ncols,
                "non_empty_cell_count": non_empty,
                "formula_count": 0,
                "formula_cached_count": 0,
                "formula_missing_count": 0,
                "formula_error_count": 0,
                "cross_sheet_formula_count": 0,
                "external_formula_count": 0,
                "direct_self_reference_count": 0,
                "formula_examples": [],
                "error_cell_count": len(error_cells),
                "error_cells": error_cells[:FORMULA_EXAMPLE_LIMIT],
                "hidden_rows": [],
                "hidden_columns": [],
                "merged_ranges": [],
                "tables": [],
                "formula_visibility": "cached_values_only",
            }
        )
    return {
        "schema_version": MANIFEST_SCHEMA_VERSION,
        "file_hash": file_hash,
        "source_name": path.name,
        "format": "xls",
        "file_size": path.stat().st_size,
        "calculation_properties": {},
        "defined_names": [],
        "sheet_count": len(sheets),
        "formula_count": 0,
        "formula_cached_count": 0,
        "formula_missing_count": 0,
        "formula_error_count": sum(sheet["formula_error_count"] for sheet in sheets),
        "has_unresolved_formulas": False,
        "formula_visibility": "cached_values_only",
        "sheets": sheets,
    }


def scan_workbook(
    path: str | os.PathLike[str],
    *,
    file_hash: str | None = None,
) -> dict[str, Any]:
    """Return a JSON-safe structural and formula-cache manifest."""
    source = Path(path).resolve()
    if not source.is_file():
        raise WorkbookScanError(f"文件不存在: {source.name}")
    extension = source.suffix.lower()
    if extension not in {".xlsx", ".xlsm", ".xls"}:
        raise WorkbookScanError(f"不支持的 Excel 格式: {extension or '无扩展名'}")
    digest = file_hash or sha256_path(source)
    if extension == ".xls":
        return _scan_xls(source, digest)
    return _scan_ooxml(source, digest)


def _atomic_copy(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=destination.parent, prefix=".copy-", delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        shutil.copy2(source, temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _atomic_json(payload: dict[str, Any], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        dir=destination.parent,
        prefix=".manifest-",
        suffix=".json",
        delete=False,
    ) as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        temporary = Path(handle.name)
    try:
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def snapshot_and_scan(
    source_path: str | os.PathLike[str],
    work_dir: str | os.PathLike[str],
) -> dict[str, Any]:
    """Create or reuse an immutable content-addressed snapshot and manifest."""
    source = Path(source_path).resolve()
    workspace = Path(work_dir).resolve()
    if not source.is_file():
        raise WorkbookScanError(f"文件不存在: {source.name}")
    digest = sha256_path(source)
    cache_dir = workbook_cache_root(workspace) / IMPORT_CACHE_DIR_NAME / digest
    snapshot = cache_dir / f"original{source.suffix.lower()}"
    manifest_path = cache_dir / "manifest.json"

    if not snapshot.is_file() or sha256_path(snapshot) != digest:
        _atomic_copy(source, snapshot)

    if manifest_path.is_file():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if (
                manifest.get("schema_version") == MANIFEST_SCHEMA_VERSION
                and manifest.get("file_hash") == digest
            ):
                return {
                    **manifest,
                    "snapshot_path": str(snapshot),
                    "manifest_path": str(manifest_path),
                    "reused_manifest": True,
                }
        except (OSError, ValueError, TypeError):
            pass

    manifest = scan_workbook(snapshot, file_hash=digest)
    manifest["source_name"] = source.name
    _atomic_json(manifest, manifest_path)
    return {
        **manifest,
        "snapshot_path": str(snapshot),
        "manifest_path": str(manifest_path),
        "reused_manifest": False,
    }


__all__ = [
    "WorkbookScanError",
    "scan_workbook",
    "snapshot_and_scan",
    "sha256_path",
]
