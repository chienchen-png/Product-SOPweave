"""Excel → SQLite project data plane importer — Phase 14A.

Scans and snapshots a workbook, verifies formula caches through an external
calculation adapter when needed, detects workbook-authored regions, validates
types and structure, then atomically materializes typed SQLite tables with
source row/column lineage.
"""

from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from storage.project_data_store import ProjectDataStore
from agent_runtime.workspace import path_lock, resolve_path
from tools.table_layout_detector import Layout, audit_and_correct, detect_layout
from tools.workbook_calculation import ensure_calculated_workbook
from tools.workbook_scan import read_workbook_values, sha256_path, snapshot_and_scan

# Per-sheet row cap: oversized sheets are rejected before any data is read.
MAX_ROWS_PER_SHEET = 200_000
IMPORT_VERSION = 8
# Sample values recorded per column in headers_json (2-3 per PRD §3.12.5).
SAMPLE_LIMIT = 3
EXCEL_ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
}


class FormulaCalculationRequired(ValueError):
    """Raised when formula values cannot be verified by an available engine."""

    def __init__(self, calculation: dict[str, Any]) -> None:
        self.calculation = calculation
        missing = int(calculation.get("missing_after") or 0)
        errors = int(calculation.get("errors_after") or 0)
        if errors or (calculation.get("stale_cache") and not missing):
            reasons = []
            if missing:
                reasons.append(f"{missing} 个公式缺少可验证的计算值")
            if errors:
                locations = [
                    f"{sheet.get('name')}!{cell}"
                    for sheet in (calculation.get("manifest") or {}).get("sheets", [])
                    for cell in sheet.get("error_cells", [])
                ][:5]
                reasons.append(
                    f"{errors} 个 Excel 错误值"
                    + ("（" + ", ".join(locations) + "）" if locations else "")
                )
            if calculation.get("stale_cache"):
                reasons.append("工作簿标记需要重新计算，缓存值尚未验证")
            super().__init__(
                "工作簿需要复核："
                + "；".join(reasons)
                + "。请在 Excel 中修复并完整重算后保存，再重新导入"
            )
            return
        super().__init__(
            f"工作簿有 {missing} 个公式缺少可验证的计算值；"
            "请安装 Microsoft Excel 或 LibreOffice 后重试，"
            "或在 Excel 中完整重算并保存文件"
        )


class WorkbookTooLargeError(ValueError):
    """Raised before reading when a sheet would otherwise be truncated."""


class NeedAnchorError(Exception):
    """Raised when a sheet's header row could not be auto-detected with
    sufficient confidence and no user anchor was provided.

    Carries the probe result (per-sheet layouts + preview rows) so the
    caller can prompt the user to confirm the header row.
    """

    def __init__(
        self,
        file_hash: str,
        layouts: dict[str, Layout],
        previews: dict[str, list[list]],
        sheet_info: list[dict[str, Any]] | None = None,
    ) -> None:
        super().__init__("需要用户确认表头行")
        self.file_hash = file_hash
        self.layouts = layouts
        self.previews = previews
        self.sheet_info = sheet_info or []


def _resolve_work_path(work_dir: str, rel_path: str) -> str:
    """Resolve a project-relative path inside the work directory."""
    return str(resolve_path(Path(work_dir), rel_path))


def sha256_file(path: str) -> str:
    """Compatibility wrapper for the shared workbook content hash."""
    return sha256_path(path)


def _sanitize_column(name: Any, used: set[str], source_col: int | None = None) -> str:
    """Sanitize a raw header cell into a safe, unique SQLite column name."""
    raw = "" if name is None else str(name)
    base = re.sub(r"[^\w\u4e00-\u9fff]+", "_", raw).strip("_")
    if not base:
        base = "Col"
    candidate = base
    if candidate in used and source_col is not None:
        from openpyxl.utils import get_column_letter

        candidate = f"{base}_{get_column_letter(source_col + 1)}"
    suffix = 1
    while candidate in used:
        suffix += 1
        candidate = f"{base}_{suffix}"
    used.add(candidate)
    return candidate


def anchor_integer(value: Any) -> int:
    """Accept integer coordinates without silently rounding floats/bools."""
    if isinstance(value, bool) or not isinstance(value, (int, str)):
        raise ValueError("导入区域的行列坐标必须是整数")
    return int(value)


def _apply_layout_override(layout: Layout, rows: list[list], value: Any) -> None:
    """Apply a user-confirmed region; API coordinates are zero based."""
    override = value if isinstance(value, dict) else {"header_row": value}
    header_row = anchor_integer(override.get("header_row", layout.header_row))
    if header_row < 0 or header_row >= len(rows):
        raise ValueError(f"表头行超出范围: {header_row + 1}")
    detected_names = {}
    if header_row == layout.header_row:
        detected_columns = [
            index
            for index in range(layout.start_col, layout.end_col + 1)
            if index not in layout.excluded_columns
        ]
        detected_names = dict(zip(detected_columns, layout.columns))
    header = rows[header_row]
    named = [
        index
        for index, cell in enumerate(header)
        if cell is not None and str(cell).strip()
    ]
    default_start = min(named) if named else 0
    default_end = max(named) if named else max(len(header) - 1, 0)
    start_col = anchor_integer(override.get("start_col", default_start))
    end_col = anchor_integer(override.get("end_col", default_end))
    width = max(map(len, rows))
    if start_col < 0 or end_col < start_col or end_col >= width:
        raise ValueError("导入列范围无效")
    if isinstance(value, dict):
        excluded = sorted(
            {
                anchor_integer(index)
                for index in (override.get("excluded_columns") or [])
            }
        )
        if any(index < start_col or index > end_col for index in excluded):
            raise ValueError("排除列必须位于导入列范围内")
    else:
        excluded = list(range(end_col + 1, len(header)))
    used: set[str] = set()
    source_columns = [
        index for index in range(start_col, end_col + 1) if index not in excluded
    ]
    if not source_columns:
        raise ValueError("导入区域至少需要保留一列")
    default_data_start = (
        layout.data_start_row if header_row == layout.header_row else header_row + 1
    )
    data_start = anchor_integer(override.get("data_start_row", default_data_start))
    data_end = anchor_integer(override.get("data_end_row", layout.data_end_row))
    if (
        data_start <= header_row
        or data_start >= len(rows)
        or data_end < -1
        or data_end >= len(rows)
        or (data_end >= 0 and data_end < data_start)
    ):
        raise ValueError("数据行范围无效")
    layout.header_row = header_row
    layout.data_start_row = data_start
    layout.data_end_row = data_end
    layout.start_col = start_col
    layout.end_col = end_col
    layout.excluded_columns = excluded
    layout.columns = [
        _sanitize_column(
            detected_names.get(index, header[index] if index < len(header) else None),
            used,
            index,
        )
        for index in source_columns
    ]
    layout.confidence = 1.0
    layout.region_source = "user_anchor"
    layout.classification = "table"


def _sqlite_type(value: Any) -> str:
    """Logical cell type. Empty values are filtered by the caller."""
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, Decimal):
        return "decimal"
    if isinstance(value, float):
        return "decimal"
    if isinstance(value, (datetime, date, time)):
        return "date"
    if isinstance(value, str) and value.strip().upper() in EXCEL_ERROR_VALUES:
        return "error"
    return "text"


def _promote_type(current: str, new: str) -> str:
    """Merge two per-column inferred types (widening)."""
    if current == new:
        return current
    if {current, new} <= {"int", "decimal"}:
        return "decimal"
    return "mixed"


def _ddl_type(logical: str) -> str:
    """Map logical type to SQLite column DDL type."""
    if logical == "int":
        return "INTEGER"
    if logical == "bool":
        return "INTEGER"
    if logical == "decimal":
        return "REAL"
    if logical == "date":
        return "TEXT"
    return "TEXT"


def _norm_value(value: Any) -> Any:
    """Normalize a cell for storage (ISO datetimes, int floats, bool→int)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


class ProjectDataImporter:
    """Materialize one Excel workbook into the project data plane."""

    def __init__(self, db) -> None:
        self.db = db
        self.store = ProjectDataStore(db)

    # ------------------------------------------------------------------
    # Public entry
    # ------------------------------------------------------------------
    def read_workbook_rows(
        self, full: str, ext: str, manifest: dict | None = None
    ) -> dict[str, list[list]]:
        """Read every sheet into an in-memory list of rows (capped).

        Returns {sheet_name: [[cell, ...], ...]}.  Used both for probing
        (detect_layout) and for materialization with a user anchor.
        """
        return read_workbook_values(
            full, max_rows=MAX_ROWS_PER_SHEET, manifest=manifest
        )

    def import_workbook(
        self,
        project_id: str,
        rel_path: str,
        work_dir: str,
        anchor_rows: dict[str, Any] | None = None,
        always_anchor: bool = False,
        progress: Callable[[int, int, str], None] | None = None,
        *,
        expected_file_hash: str | None = None,
        calculation_strategy: str = "auto",
    ) -> dict:
        """Serialize imports of one canonical source, including aliases."""
        full = Path(_resolve_work_path(work_dir, rel_path))
        canonical = full.relative_to(Path(work_dir).resolve()).as_posix()
        with path_lock(full):
            return self._import_workbook(
                project_id,
                canonical,
                work_dir,
                anchor_rows,
                always_anchor,
                progress,
                expected_file_hash=expected_file_hash,
                calculation_strategy=calculation_strategy,
            )

    def _import_workbook(
        self,
        project_id: str,
        rel_path: str,
        work_dir: str,
        anchor_rows: dict[str, Any] | None = None,
        always_anchor: bool = False,
        progress: Callable[[int, int, str], None] | None = None,
        *,
        expected_file_hash: str | None = None,
        calculation_strategy: str = "auto",
    ) -> dict:
        """Import a workbook.  Returns {"file_id", "reused", "file", "sheets"}.

        anchor_rows: {sheet_name: header_row} confirmed by the user.
        When a sheet's auto-detected confidence is low and no anchor is
        given for it, NeedAnchorError is raised with the probe result.

        always_anchor: True 时无论置信度高低都要求用户确认表头行
        （应用「每次上传都确认变量行/数据行」场景）。
        """
        full = _resolve_work_path(work_dir, rel_path)
        if not os.path.isfile(full):
            raise ValueError(f"文件不存在: {rel_path}")
        ext = os.path.splitext(full)[1].lower()
        if ext == ".xlsm":
            raise ValueError(
                "数据库导入暂不支持带宏的 .xlsm 文件，请另存为 .xlsx 后导入"
            )
        if ext not in {".xlsx", ".xls"}:
            raise ValueError(f"仅支持 Excel 文件: {rel_path}")

        scan = snapshot_and_scan(full, work_dir)
        file_hash = scan["file_hash"]
        if expected_file_hash and file_hash != expected_file_hash:
            raise ValueError("源文件已变化，请重新导入并确认新的导入区域")

        # A source is identified by project + relative path. Content hashes
        # share immutable snapshots, never writable datasets or writeback paths.
        existing = self.store.find_file_by_source(project_id, rel_path)
        if (
            existing
            and existing.get("file_hash") == file_hash
            and existing.get("status") == "done"
            and not always_anchor
            and not anchor_rows
        ):
            try:
                existing_meta = json.loads(existing.get("meta_json") or "{}")
            except (TypeError, ValueError):
                existing_meta = {}
            if int(existing_meta.get("import_version") or 0) >= IMPORT_VERSION and bool(
                existing_meta.get("formulas_materialized")
            ):
                return {
                    "file_id": existing["id"],
                    "reused": True,
                    "file": existing,
                    "sheets": self.store.list_sheets(existing["id"]),
                }

        calculation_strategy = str(calculation_strategy or "auto").strip().lower()
        if calculation_strategy not in {"auto", "cached"}:
            raise ValueError("不支持的公式处理策略")
        calculation = ensure_calculated_workbook(
            scan, allow_stale_cache=calculation_strategy == "cached"
        )
        if calculation.get("requires_review"):
            raise FormulaCalculationRequired(calculation)
        read_source = str(calculation["calculated_path"])
        effective_manifest = calculation.get("manifest") or scan
        oversized = []
        for sheet in effective_manifest.get("sheets") or []:
            bounds = sheet.get("actual_bounds") or []
            if len(bounds) == 4 and int(bounds[2] or 0) > MAX_ROWS_PER_SHEET:
                oversized.append(f'{sheet.get("name")}: {bounds[2]} 行')
        if oversized:
            raise WorkbookTooLargeError(
                f"工作簿超过单表 {MAX_ROWS_PER_SHEET} 行上限，已阻止截断导入："
                + "；".join(oversized)
            )

        # ── Read all rows once, probe layouts, apply anchors ──
        sheets_rows = self.read_workbook_rows(read_source, ext, effective_manifest)
        scan_sheets = {
            item.get("name"): item for item in (effective_manifest.get("sheets") or [])
        }
        for defined_name in effective_manifest.get("defined_names") or []:
            for destination in defined_name.get("destinations") or []:
                sheet_meta = scan_sheets.get(destination.get("sheet"))
                if sheet_meta is not None:
                    sheet_meta.setdefault("named_ranges", []).append(
                        {
                            "name": defined_name.get("name"),
                            "ref": destination.get("ref"),
                        }
                    )
        layouts = {
            name: detect_layout(rows, scan_sheets.get(name), name)
            for name, rows in sheets_rows.items()
        }
        anchor_rows = anchor_rows or {}
        for name, hr in anchor_rows.items():
            if name == "__all__":
                # 通配锚点：应用到所有低置信 sheet（LLM 从用户处获取的
                # 通用表头行号，多数财务文件共用同一布局）。
                for nm, lay in list(layouts.items()):
                    if lay.is_confident or nm not in sheets_rows:
                        continue
                    _apply_layout_override(lay, sheets_rows[nm], hr)
                continue
            if name in layouts and name in sheets_rows:
                _apply_layout_override(layouts[name], sheets_rows[name], hr)

        # ── Col audit & self-correction (Phase 14E #2) ──
        # Blank header cells become "Col" columns.  Try to recover real
        # names from a structurally isomorphic sheet.
        # Unrecovered Cols are reported so the caller can warn the user.
        uncorrected = audit_and_correct(sheets_rows, layouts)
        for name, indices in uncorrected.items():
            layout = layouts[name]
            source_columns = [
                index
                for index in range(layout.start_col, layout.end_col + 1)
                if index not in layout.excluded_columns
            ]
            uncorrected[name] = [source_columns[index] for index in indices]
        uncorrected = {name: idxs for name, idxs in uncorrected.items() if idxs}

        # ── 需要用户确认表头的 sheet ──
        # 默认：低置信度才问；always_anchor=True：全部 sheet 都问（每次上传确认）
        if always_anchor:
            low = {
                name: lay
                for name, lay in layouts.items()
                if name in sheets_rows
                and name not in anchor_rows
                and lay.classification not in {"empty", "notes"}
            }
        else:
            low = {
                name: lay
                for name, lay in layouts.items()
                if not lay.is_confident
                and name not in anchor_rows
                and lay.classification not in {"empty", "notes"}
            }
        if low:
            raise NeedAnchorError(
                file_hash=file_hash,
                layouts=layouts,
                previews={name: sheets_rows[name][:8] for name in low},
                sheet_info=[
                    {
                        "name": name,
                        "state": info.get("state", "visible"),
                        "hidden_rows": (info.get("hidden_rows") or [])[:20],
                        "hidden_row_count": len(info.get("hidden_rows") or []),
                        "hidden_columns": (info.get("hidden_columns") or [])[:20],
                        "hidden_column_count": len(info.get("hidden_columns") or []),
                    }
                    for name, info in scan_sheets.items()
                ],
            )

        file_name = os.path.basename(full)
        file_rec = self.store.create_file(
            project_id, rel_path, file_name, file_hash, ext.lstrip(".")
        )
        file_id = file_rec["id"]
        try:
            sheets_meta: list[dict] = []
            importable = [
                (name, rows)
                for name, rows in sheets_rows.items()
                if layouts[name].classification not in {"empty", "notes"}
            ]
            if not importable:
                raise ValueError("工作簿没有可导入的数据区域")
            self.db.execute("BEGIN")
            total = len(importable)
            for idx, (name, rows) in enumerate(importable):
                meta = self._materialize_sheet(file_id, idx, name, rows, layouts[name])
                sheets_meta.append(meta)
                if progress:
                    progress(idx + 1, total, name)
        except Exception as exc:
            self.db.rollback()
            self.store.update_file(
                file_id, status="failed", error_message=str(exc)[:500]
            )
            self.store.cleanup_sheets(file_id)
            raise

        total_rows = sum(s["row_count"] for s in sheets_meta)
        layout_meta = {
            name: {
                "header_row": lay.header_row,
                "data_start_row": lay.data_start_row,
                "data_end_row": lay.data_end_row,
                "start_col": lay.start_col,
                "end_col": lay.end_col,
                "excluded_columns": lay.excluded_columns,
                "classification": lay.classification,
                "region_source": lay.region_source,
                "confidence": lay.confidence,
                "anchored_by": "user" if name in anchor_rows else "auto",
            }
            for name, lay in layouts.items()
        }
        snapshot_relpath = os.path.relpath(str(scan["snapshot_path"]), work_dir)
        calculated_relpath = os.path.relpath(read_source, work_dir)
        workbook_manifest = {
            key: value
            for key, value in scan.items()
            if key not in {"snapshot_path", "manifest_path", "reused_manifest"}
        }
        calculation_meta = {
            key: value
            for key, value in calculation.items()
            if key not in {"calculated_path", "manifest"}
        }
        meta = {
            "import_version": IMPORT_VERSION,
            "file_size": os.path.getsize(full),
            "snapshot_relpath": snapshot_relpath,
            "calculated_relpath": calculated_relpath,
            "formulas_materialized": not bool(
                effective_manifest.get("has_unresolved_formulas")
            ),
            "formula_audit": {
                "formula_count": int(scan.get("formula_count") or 0),
                "cached_count": int(
                    effective_manifest.get("formula_cached_count") or 0
                ),
                "missing_count": int(
                    effective_manifest.get("formula_missing_count") or 0
                ),
                "error_count": int(effective_manifest.get("formula_error_count") or 0),
                "engine": calculation.get("engine"),
            },
            "calculation": calculation_meta,
            "workbook_manifest": workbook_manifest,
            "layouts": layout_meta,
            "validation": {
                "status": (
                    "warning"
                    if any(item["validation"]["issues"] for item in sheets_meta)
                    else "ok"
                ),
                "sheets": {
                    item["sheet_name"]: item["validation"] for item in sheets_meta
                },
                "skipped_sheets": [
                    {"name": name, "classification": lay.classification}
                    for name, lay in layouts.items()
                    if lay.classification in {"empty", "notes"}
                ],
            },
            "uncorrected_cols": uncorrected,
            "col_audit": "ok" if not uncorrected else "has_unnamed",
        }
        try:
            self.store.update_file(
                file_id,
                status="done",
                error_message=None,
                sheet_count=len(sheets_meta),
                total_rows=total_rows,
                meta_json=json.dumps(meta, ensure_ascii=False),
                commit=False,
            )
            if existing:
                self.store.delete_file(existing["id"], commit=False)
            self.db.commit()
        except Exception as exc:
            self.db.rollback()
            self.store.update_file(
                file_id, status="failed", error_message=str(exc)[:500]
            )
            raise
        return {
            "file_id": file_id,
            "reused": False,
            "file": self.store.get_file(file_id),
            "sheets": sheets_meta,
        }

    # ------------------------------------------------------------------
    # Sheet materialization (anchor-aware)
    # ------------------------------------------------------------------
    def _materialize_sheet(
        self,
        file_id: str,
        sheet_index: int,
        sheet_name: str,
        rows: list[list],
        layout: Layout,
    ) -> dict:
        table_name = f"data_{file_id[:8]}_{sheet_index}"
        header_row = layout.header_row
        data_start = layout.data_start_row
        data_end = layout.data_end_row

        if header_row >= len(rows):
            raise ValueError(f"工作表 {sheet_name} 表头超出范围")

        start_col = max(0, int(layout.start_col or 0))
        end_col = int(layout.end_col)
        if end_col < start_col:
            end_col = len(rows[header_row]) - 1

        source_columns = [
            index
            for index in range(start_col, end_col + 1)
            if index not in set(layout.excluded_columns)
        ]
        raw_names = [
            rows[header_row][index] if index < len(rows[header_row]) else None
            for index in source_columns
        ]
        used: set[str] = {"id", "_source_row", "_source_region"}
        # Prefer the detector's (possibly self-corrected) column names; fall
        # back to re-sanitizing the raw header row for sheets materialized
        # with a user anchor (layout.columns may be empty there).
        if layout.columns:
            columns = [
                _sanitize_column(c, used, source_columns[offset])
                for offset, c in enumerate(layout.columns)
            ]
        else:
            columns = [
                _sanitize_column(
                    rows[header_row][index] if index < len(rows[header_row]) else None,
                    used,
                    index,
                )
                for index in source_columns
            ]
        col_types: list[str | None] = [None] * len(columns)
        col_type_counts: list[dict[str, int]] = [{} for _ in columns]
        if not columns or all(name == "Col" for name in columns):
            raise ValueError(f"工作表 {sheet_name} 表头为空")

        # Data region: data_start .. data_end (data_end inclusive, -1 → end)
        end = data_end
        if end < 0 or end >= len(rows):
            end = len(rows) - 1

        data_rows: list[list] = []
        source_rows: list[int] = []
        for idx in range(data_start, end + 1):
            if len(data_rows) >= MAX_ROWS_PER_SHEET:
                break
            raw = rows[idx]
            row = [raw[index] if index < len(raw) else None for index in source_columns]
            while len(row) < len(columns):
                row.append(None)
            # Phase 17 P6：跳过数据区内的全空行（财务表常用空行做视觉分隔），
            # 避免产生全 null 数据行；不改变行号锚定（写回仍按行序）。
            if all(v is None or v == "" for v in row):
                continue
            for i, v in enumerate(row):
                if v is not None and v != "":
                    t = _sqlite_type(v)
                    col_type_counts[i][t] = col_type_counts[i].get(t, 0) + 1
                    col_types[i] = (
                        t if col_types[i] is None else _promote_type(col_types[i], t)
                    )
            data_rows.append([_norm_value(v) for v in row])
            source_rows.append(idx + 1)

        # Columns with no non-null values fall back to text.
        col_types = [t or "text" for t in col_types]

        col_defs = [
            {"name": columns[i], "type": _ddl_type(col_types[i])}
            for i in range(len(columns))
        ]
        internal_defs = [
            {"name": "_source_row", "type": "INTEGER"},
            {"name": "_source_region", "type": "TEXT"},
        ]
        self.store.create_data_table(table_name, internal_defs + col_defs, commit=False)
        from openpyxl.utils import get_column_letter

        region = (
            f"{sheet_name}!{get_column_letter(start_col + 1)}{data_start + 1}:"
            f"{get_column_letter(end_col + 1)}{end + 1}"
        )
        stored_rows = [
            [source_rows[index], region, *row] for index, row in enumerate(data_rows)
        ]
        self.store.insert_rows(
            table_name,
            ["_source_row", "_source_region", *columns],
            stored_rows,
            commit=False,
        )

        headers = self._build_headers(
            columns,
            col_types,
            data_rows,
            source_columns=source_columns,
            original_names=raw_names,
            type_counts=col_type_counts,
        )
        self.store.create_sheet(
            file_id,
            sheet_name,
            sheet_index,
            table_name,
            len(data_rows),
            len(columns),
            headers,
            commit=False,
        )
        validation = self._validation_report(rows, layout, raw_names, columns, headers)
        return {
            "sheet_name": sheet_name,
            "sheet_index": sheet_index,
            "table_name": table_name,
            "row_count": len(data_rows),
            "col_count": len(columns),
            "validation": validation,
        }

    def _build_headers(
        self,
        columns: list[str],
        col_types: list[str],
        data_rows: list[list],
        source_columns: list[int] | None = None,
        original_names: list[Any] | None = None,
        type_counts: list[dict[str, int]] | None = None,
    ) -> list[dict]:
        n = len(data_rows)
        headers: list[dict] = []
        for i, name in enumerate(columns):
            samples: list[str] = []
            non_null = 0
            counts = type_counts[i] if type_counts else {}
            for row in data_rows:
                v = row[i] if i < len(row) else None
                if v is not None and v != "":
                    non_null += 1
                    if len(samples) < SAMPLE_LIMIT:
                        samples.append(str(v))
            source_col = source_columns[i] if source_columns else i
            from openpyxl.utils import get_column_letter

            headers.append(
                {
                    "name": name,
                    "original_name": (
                        ""
                        if not original_names or original_names[i] is None
                        else str(original_names[i])
                    ),
                    "source_col": source_col,
                    "excel_col": get_column_letter(source_col + 1),
                    "type": col_types[i],
                    "type_counts": counts,
                    "type_confidence": (
                        round(max(counts.values()) / non_null, 4) if non_null else 1.0
                    ),
                    "non_null_rate": round(non_null / n, 4) if n else 1.0,
                    "samples": samples,
                }
            )
        return headers

    def _validation_report(
        self,
        rows: list[list],
        layout: Layout,
        raw_names: list[Any],
        columns: list[str],
        headers: list[dict],
    ) -> dict[str, Any]:
        """Build deterministic, coordinate-aware import diagnostics."""
        issues: list[dict[str, Any]] = []
        normalized_names = [str(value or "").strip() for value in raw_names]
        blanks = [
            headers[index]["source_col"]
            for index, value in enumerate(normalized_names)
            if not value
        ]
        if blanks:
            issues.append({"code": "blank_headers", "columns": blanks})
        duplicates = sorted(
            {
                value
                for value in normalized_names
                if value and normalized_names.count(value) > 1
            }
        )
        if duplicates:
            issues.append({"code": "duplicate_headers", "names": duplicates})
        for header in headers:
            if header["type"] in {"mixed", "error"}:
                issues.append(
                    {
                        "code": "mixed_or_error_type",
                        "column": header["name"],
                        "type_counts": header["type_counts"],
                    }
                )
            if header["non_null_rate"] < 0.2:
                issues.append(
                    {
                        "code": "high_null_rate",
                        "column": header["name"],
                        "non_null_rate": header["non_null_rate"],
                    }
                )
        missing_rows: list[int] = []
        extra_cells: list[str] = []
        excluded_count = 0
        included_columns = {header["source_col"] for header in headers}
        start = max(layout.data_start_row, 0)
        end = layout.data_end_row if layout.data_end_row >= 0 else len(rows) - 1
        from openpyxl.utils import get_column_letter

        for row_index in range(start, len(rows)):
            row = rows[row_index]
            if row_index <= end and len(row) <= layout.end_col:
                missing_rows.append(row_index + 1)
            for col_index, value in enumerate(row):
                if value in (None, "") or (row_index <= end and col_index in included_columns):
                    continue
                excluded_count += 1
                if len(extra_cells) < 20:
                    extra_cells.append(f"{get_column_letter(col_index + 1)}{row_index + 1}")
        if missing_rows:
            issues.append({"code": "short_rows", "rows": missing_rows[:20]})
        if extra_cells:
            issues.append({"code": "excluded_non_empty_cells", "cells": extra_cells, "count": excluded_count})
        return {
            "status": "warning" if issues else "ok",
            "row_count": sum(
                1
                for row in rows[start : end + 1]
                if any(value not in (None, "") for value in row)
            ),
            "column_count": len(columns),
            "issues": issues,
        }
