"""Generate small synthetic Office examples; no business data is copied."""
from pathlib import Path
import argparse


def create_demo(directory: Path) -> None:
    import openpyxl
    import xlwt
    from docx import Document

    directory.mkdir(parents=True, exist_ok=True)
    rows = [("编号", "部门", "收入", "成本"), ("001", "研发", 120, 80), ("002", "销售", 150, 90)]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "明细"
    for row in rows:
        sheet.append(row)
    book.save(directory / "人工样例.xlsx")
    book.close()
    old = xlwt.Workbook()
    sheet = old.add_sheet("明细")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    old.save(str(directory / "人工样例.xls"))
    document = Document()
    document.add_heading("人工样例说明", 0)
    document.add_paragraph("收入合计 270，成本合计 170，利润合计 100。编号 001/002 必须保留前导零。")
    document.save(directory / "人工样例.docx")
    (directory / "报告.md").write_text("# 人工样例报告\n\n收入 **270**，成本 **170**，利润 **100**。\n", encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    create_demo(parser.parse_args().output)
