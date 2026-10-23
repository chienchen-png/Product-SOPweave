import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from pathlib import Path

from flask import Flask


class AppServicesTests(unittest.TestCase):
    def test_public_route_contract_is_complete(self):
        from routes.apps import apps_bp, finmod_bp
        app = Flask(__name__)
        app.register_blueprint(apps_bp)
        app.register_blueprint(finmod_bp)
        urls = {rule.rule for rule in app.url_map.iter_rules() if rule.rule.startswith('/api/apps')}
        expected = {f'/api/apps/recon/{suffix}' for suffix in ['verify', 'writeback', 'audit']}
        expected.update(f'/api/apps/finmod/{suffix}' for suffix in ['meta', 'models/<model_id>', 'charts/<chart_id>', 'guide', 'varmap', 'varmap-suggest', 'prep-data', 'derived-eval', 'charts-suggest', 'run', 'analyze', 'py-chart', 'save-chart', 'files/<run_id>/assets/<path:filepath>', 'export-md', 'export-docx'])
        self.assertEqual(urls, expected)
        bad = app.test_client().post('/api/apps/finmod/run', json={'project_id': ['invalid'], 'model_ids': 'invalid'})
        self.assertEqual(bad.status_code, 400)

    def test_reconciliation_and_finmod_report_prompts_have_separate_owners(self):
        from services import reconciliation, finmod_reports
        self.assertIn('数据核对与校验', reconciliation.REPORT_SYSTEM)
        self.assertNotEqual(reconciliation.REPORT_SYSTEM, finmod_reports.REPORT_SYSTEM)

    def test_report_directory_survives_restart_and_parallel_allocation(self):
        from storage.report_store import ReportStore
        with tempfile.TemporaryDirectory(prefix='报告 并发 ') as directory:
            root = Path(directory)
            path = root / 'app.db'
            def connect():
                db = sqlite3.connect(path, timeout=10)
                db.row_factory = sqlite3.Row
                return db
            with closing(connect()) as db:
                db.executescript(Path('storage/schema.sql').read_text(encoding='utf-8'))
                from storage.project_store import ProjectStore
                project = ProjectStore(db).create('report', str(root))
                first = ReportStore(db).allocate(project['id'], 'run', root)
                (first / 'keep.md').write_text('edited', encoding='utf-8')
            with closing(connect()) as db:
                again = ReportStore(db).allocate(project['id'], 'run', root)
                self.assertEqual(again, first)
                self.assertEqual((again / 'keep.md').read_text(encoding='utf-8'), 'edited')
            def allocate(index):
                with closing(connect()) as db:
                    return ReportStore(db).allocate(project['id'], f'run-{index}', root)
            with ThreadPoolExecutor(max_workers=4) as pool:
                folders = list(pool.map(allocate, range(8)))
            self.assertEqual(len(set(folders)), 8)
            self.assertNotIn(first, folders)

    def test_generating_same_report_cannot_overwrite_user_edits(self):
        from services.finmod_reports import finmod_export_md
        from services.results import ServiceError
        from storage.connections import connection_scope
        from storage.project_store import ProjectStore
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            def connect():
                db = sqlite3.connect(root / 'app.db')
                db.row_factory = sqlite3.Row
                return db
            with connection_scope(connect) as db:
                db.executescript(Path('storage/schema.sql').read_text(encoding='utf-8'))
                project = ProjectStore(db).create('report', str(root))
                command = {'project_id': project['id'], 'run_id': 'r', 'model_id': 'fake'}
                with patch('services.finmod_reports._build_llm_client', return_value=object()), patch('services.finmod_reports._chat_complete', return_value='# report'):
                    outcome = finmod_export_md(command)
                    report = Path(outcome.data['path'])
                    report.write_text('# user edit', encoding='utf-8')
                    with self.assertRaises(ServiceError) as caught:
                        finmod_export_md(command)
                self.assertEqual(caught.exception.status_code, 409)
                self.assertEqual(report.read_text(encoding='utf-8'), '# user edit')

    def test_reconciliation_writeback_rejects_source_changed_after_import(self):
        import hashlib

        import openpyxl

        from services.reconciliation import _writeback_export
        from storage.project_data_store import ProjectDataStore

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            database = sqlite3.connect(root / 'app.db')
            database.row_factory = sqlite3.Row
            try:
                database.executescript(
                    Path('storage/schema.sql').read_text(encoding='utf-8')
                )
                project_id = 'project-recon-stale'
                workbook_path = root / '分表.xlsx'
                workbook = openpyxl.Workbook()
                workbook.active.append(['编号', '金额'])
                workbook.active.append(['C001', 100])
                workbook.save(workbook_path)
                imported_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

                workbook = openpyxl.load_workbook(workbook_path)
                workbook.active['B2'] = 999
                workbook.save(workbook_path)
                externally_changed_bytes = workbook_path.read_bytes()
                externally_changed_hash = hashlib.sha256(
                    externally_changed_bytes
                ).hexdigest()

                database.execute(
                    """
                    INSERT INTO project_data_files(
                        id,project_id,source_path,file_name,file_hash,format,
                        sheet_count,total_rows,status,meta_json,created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        'file-recon-stale',
                        project_id,
                        '分表.xlsx',
                        '分表.xlsx',
                        imported_hash,
                        'xlsx',
                        1,
                        1,
                        'done',
                        '{"layouts":{}}',
                        '2026-10-21T00:00:00+00:00',
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                database.commit()

                result = _writeback_export(
                    ProjectDataStore(database), project_id, '分表.xlsx', str(root)
                )

                self.assertFalse(result['success'])
                self.assertIn('源文件已更改', result['error'])
                self.assertEqual(
                    hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                    externally_changed_hash,
                )
                self.assertEqual(workbook_path.read_bytes(), externally_changed_bytes)
            finally:
                database.close()

    def test_reconciliation_writeback_stops_before_temp_changes_when_source_is_stale(self):
        import hashlib

        import openpyxl

        from services.reconciliation import recon_writeback
        from services.results import ServiceError

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            main_db = sqlite3.connect(root / 'main.db')
            main_db.row_factory = sqlite3.Row
            temp_db = sqlite3.connect(root / 'tmp.db')
            temp_db.row_factory = sqlite3.Row
            try:
                main_db.executescript(
                    Path('storage/schema.sql').read_text(encoding='utf-8')
                )
                for table in ('project_data_files', 'project_data_changes'):
                    statement = main_db.execute(
                        "SELECT sql FROM sqlite_master WHERE type='table' AND name=?",
                        (table,),
                    ).fetchone()[0]
                    temp_db.execute(statement)
                temp_db.commit()

                workbook_path = root / '分表.xlsx'
                workbook = openpyxl.Workbook()
                workbook.active.append(['编号', '金额'])
                workbook.active.append(['C001', 100])
                workbook.save(workbook_path)
                imported_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
                workbook = openpyxl.load_workbook(workbook_path)
                workbook.active['B2'] = 999
                workbook.save(workbook_path)
                externally_changed_bytes = workbook_path.read_bytes()

                main_db.execute(
                    """
                    INSERT INTO project_data_files(
                        id,project_id,source_path,file_name,file_hash,format,
                        sheet_count,total_rows,status,meta_json,created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        'file-recon-stale',
                        'project-recon-stale',
                        '分表.xlsx',
                        '分表.xlsx',
                        imported_hash,
                        'xlsx',
                        1,
                        1,
                        'done',
                        '{"layouts":{}}',
                        '2026-10-21T00:00:00+00:00',
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                main_db.commit()
                run = SimpleNamespace(
                    db=temp_db,
                    run_id='run-recon-stale',
                    commit=lambda: temp_db.commit(),
                )
                params = {
                    'detailFile': '分表.xlsx',
                    'baseFile': '依据表.xlsx',
                    'matchKeys': {'k1': '编号'},
                    'variables': ['金额'],
                    'baseAmountCol': '金额',
                }

                with (
                    patch('services.reconciliation.get_db', return_value=main_db),
                    patch('services.reconciliation.get_run', return_value=run),
                    patch(
                        'services.reconciliation._project',
                        return_value={'work_dir': str(root)},
                    ),
                ):
                    with self.assertRaisesRegex(ServiceError, '源文件已更改'):
                        recon_writeback(
                            {
                                'run_id': run.run_id,
                                'project_id': 'project-recon-stale',
                                'params': params,
                                'diffs': [
                                    {
                                        'level': 'A2',
                                        'contract': 'C001',
                                        'base_amount': 100,
                                    }
                                ],
                            }
                        )

                self.assertEqual(workbook_path.read_bytes(), externally_changed_bytes)
                self.assertEqual(
                    temp_db.execute(
                        'SELECT COUNT(*) FROM project_data_changes'
                    ).fetchone()[0],
                    0,
                )
            finally:
                temp_db.close()
                main_db.close()

    def test_reconciliation_writeback_rejects_changed_data_revision(self):
        import hashlib

        import openpyxl

        from services.reconciliation import recon_writeback
        from services.results import ServiceError

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            main_db = sqlite3.connect(root / 'main.db')
            main_db.row_factory = sqlite3.Row
            temp_db = sqlite3.connect(root / 'tmp.db')
            temp_db.row_factory = sqlite3.Row
            try:
                main_db.executescript(
                    Path('storage/schema.sql').read_text(encoding='utf-8')
                )
                for table in ('project_data_files', 'project_data_changes'):
                    statement = main_db.execute(
                        "SELECT sql FROM sqlite_master WHERE type='table' AND name=?",
                        (table,),
                    ).fetchone()[0]
                    temp_db.execute(statement)
                temp_db.commit()

                workbook_path = root / '分表.xlsx'
                workbook = openpyxl.Workbook()
                workbook.active.append(['编号', '金额'])
                workbook.active.append(['C001', 100])
                workbook.save(workbook_path)
                file_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
                main_db.execute(
                    """
                    INSERT INTO projects(
                        id,name,work_dir,work_mode,created_at,updated_at
                    ) VALUES(?,?,?,?,?,?)
                    """,
                    (
                        'project-recon-version',
                        'revision test',
                        str(root),
                        'manual',
                        '2026-10-21T00:00:00+00:00',
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                main_db.execute(
                    """
                    INSERT INTO project_data_files(
                        id,project_id,source_path,file_name,file_hash,format,
                        sheet_count,total_rows,status,meta_json,created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        'file-recon-version',
                        'project-recon-version',
                        '分表.xlsx',
                        '分表.xlsx',
                        file_hash,
                        'xlsx',
                        1,
                        1,
                        'done',
                        '{"layouts":{}}',
                        '2026-10-21T00:00:00+00:00',
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                main_db.execute(
                    """
                    INSERT INTO runs(
                        id,project_id,agent_id,model_id,title,status,
                        created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?)
                    """,
                    (
                        'run-recon-version',
                        'project-recon-version',
                        'app:recon',
                        'model',
                        '数据核对',
                        'completed',
                        '2026-10-21T00:00:00+00:00',
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                main_db.execute(
                    """
                    INSERT INTO run_events(run_id,seq,event_type,payload_json,created_at)
                    VALUES(?,?,?,?,?)
                    """,
                    (
                        'run-recon-version',
                        1,
                        'app.started',
                        json.dumps(
                            {'owner_pid': 123, 'data_revision': 'stale-revision'}
                        ),
                        '2026-10-21T00:00:00+00:00',
                    ),
                )
                main_db.commit()
                run = SimpleNamespace(
                    db=temp_db,
                    run_id='run-recon-version',
                    commit=lambda: temp_db.commit(),
                )

                with (
                    patch('services.reconciliation.get_db', return_value=main_db),
                    patch('services.reconciliation.get_run', return_value=run),
                    patch('services.app_data.get_db', return_value=main_db),
                ):
                    with self.assertRaisesRegex(ServiceError, '数据版本已变化'):
                        recon_writeback(
                            {
                                'run_id': run.run_id,
                                'project_id': 'project-recon-version',
                                'params': {
                                    'detailFile': '分表.xlsx',
                                    'baseFile': '依据表.xlsx',
                                    'matchKeys': {'k1': '编号'},
                                    'variables': ['金额'],
                                    'baseAmountCol': '金额',
                                },
                                'diffs': [],
                            }
                        )

                self.assertEqual(
                    temp_db.execute(
                        'SELECT COUNT(*) FROM project_data_changes'
                    ).fetchone()[0],
                    0,
                )
            finally:
                temp_db.close()
                main_db.close()

    def test_reconciliation_verification_degrades_when_model_is_unavailable(self):
        from services.reconciliation import recon_verify

        class FakeWorkbenchStore:
            def __init__(self, database):
                self.database = database

            def start_result(self, *args):
                return None

            def record_result(self, *args):
                return None

            def fail_result(self, *args):
                return None

        fake_run = SimpleNamespace(
            run_id='run-recon-fallback',
            db=object(),
            cleanup=lambda: None,
        )
        deterministic_result = {
            'success': True,
            'summary': {
                'total': 1,
                'matched': 1,
                'ok': 1,
                'a1': 0,
                'a2': 0,
                'b': 0,
                'c': 0,
            },
            'differences': [],
            'unmatched': [],
            'audit_items': [],
            'amount_col': '金额',
            'match_keys': ['编号'],
        }
        with (
            patch('services.reconciliation.get_db', return_value=object()),
            patch(
                'services.reconciliation._resolve_verify_params',
                return_value={
                    'base_file': '依据表.xlsx',
                    'detail_file': '分表.xlsx',
                    'base_sheet': '数据',
                    'match_keys': ['编号'],
                    'amount_col': '金额',
                    'tolerance': 1.0,
                    'unmatched_action': 'list',
                    'src_col': '金额',
                },
            ),
            patch(
                'services.reconciliation._build_llm_client',
                side_effect=ValueError('所选模型不可用或已删除'),
            ),
            patch(
                'services.reconciliation.TmpRun.create', return_value=fake_run
            ),
            patch(
                'services.reconciliation.register_run', return_value=None
            ),
            patch(
                'services.reconciliation.unregister_run', return_value=None
            ),
            patch(
                'services.reconciliation._copy_project_data_to_tmp',
                return_value=None,
            ),
            patch(
                'services.reconciliation.reconcile_variables',
                return_value=deterministic_result,
            ),
            patch('services.reconciliation.WorkbenchStore', FakeWorkbenchStore),
            patch(
                'services.workbench.persist_report',
                return_value={'rel_path': '报告/核对报告.md'},
            ),
        ):
            outcome = recon_verify(
                {
                    'project_id': 'project-recon-fallback',
                    'model_id': 'deleted-model',
                    'params': {
                        'baseFile': '依据表.xlsx',
                        'detailFile': '分表.xlsx',
                        'matchKeys': {'k1': '编号'},
                        'variables': ['金额'],
                        'baseAmountCol': '金额',
                    },
                }
            )

        self.assertEqual(outcome.data['reportMode'], 'deterministic_fallback')
        self.assertIn('当前模型不可用', outcome.data['report'])
        self.assertEqual(outcome.data['reconData']['summary']['ok'], 1)

    def test_financial_model_guide_offers_manual_path_when_model_is_unavailable(self):
        from services.finmod_guidance import _finmod_guide

        with (
            patch('services.finmod_guidance.get_db', return_value=object()),
            patch(
                'services.finmod_guidance._build_files_index',
                return_value=[],
            ),
            patch(
                'services.finmod_guidance._build_llm_client',
                side_effect=ValueError('所选模型不可用或已删除'),
            ),
        ):
            result = _finmod_guide(
                'project-finmod-fallback',
                'deleted-model',
                '分析利润变化',
            )

        self.assertFalse(result['success'])
        self.assertEqual(result['mode'], 'A')
        self.assertIn('手动', result['message'])
        self.assertEqual(result['directions'], [])

    def test_financial_model_varmap_degrades_to_deterministic_mapping(self):
        from services.finmod_guidance import finmod_varmap_suggest

        with (
            patch('services.finmod_guidance.get_db', return_value=object()),
            patch(
                'services.finmod_guidance._build_varmap_user_content',
                return_value=('columns', {'models': [], 'all_cols': []}),
            ),
            patch(
                'services.finmod_guidance._build_llm_client',
                side_effect=ValueError('所选模型不可用或已删除'),
            ),
            patch(
                'services.finmod_guidance.ProjectDataStore',
                return_value=object(),
            ),
            patch(
                'services.finmod_guidance._sheet_columns',
                return_value=[],
            ),
            patch(
                'services.finmod_guidance._numeric_candidate_cols',
                return_value=[],
            ),
        ):
            outcome = finmod_varmap_suggest(
                {
                    'project_id': 'project-finmod-fallback',
                    'file_name': '数据.xlsx',
                    'model_ids': ['g1-descriptive-stats'],
                    'model_id': 'deleted-model',
                }
            )

        self.assertEqual(outcome.status_code, 200)
        self.assertFalse(outcome.data['used_llm'])
        self.assertEqual(
            outcome.data['ai_suggestions'],
            {'mappings': [], 'stat_bindings': []},
        )

    def test_financial_model_report_degrades_without_losing_computed_result(self):
        from services.finmod_reports import finmod_export_md

        class EmptyQuery:
            def fetchone(self):
                return None

        class EmptyDatabase:
            def execute(self, *args):
                return EmptyQuery()

        with tempfile.TemporaryDirectory() as directory:
            with (
                patch('services.finmod_reports.get_db', return_value=EmptyDatabase()),
                patch(
                    'services.finmod_reports._finmod_report_dir',
                    return_value=str(Path(directory) / '财务分析报告1'),
                ),
                patch(
                    'services.finmod_reports._finmod_report_rel_base',
                    return_value='财务分析报告/财务分析报告1',
                ),
                patch(
                    'services.finmod_reports._build_llm_client',
                    side_effect=ValueError('所选模型不可用或已删除'),
                ),
            ):
                outcome = finmod_export_md(
                    {
                        'project_id': 'project-finmod-fallback',
                        'run_id': 'run-finmod-fallback',
                        'goal': '分析利润变化',
                        'model_id': 'deleted-model',
                        'models': [{'name': '描述统计', 'result': {'ok': True}}],
                        'charts': [],
                    }
                )

        self.assertFalse(outcome.data['used_llm'])
        self.assertIn('确定性', outcome.data['report'])
        self.assertTrue(outcome.data['report'].strip())

    def test_financial_model_period_selection_reaches_each_model_parameters(self):
        from services.finmod_compute import _model_params_with_period

        result = _model_params_with_period(
            {'d4-yoy-mom': {'top_n': 3}, 'g1-descriptive-stats': {}},
            ['d4-yoy-mom', 'g1-descriptive-stats'],
            '日期',
        )

        self.assertEqual(
            result,
            {
                'd4-yoy-mom': {'top_n': 3, 'period_col': '日期'},
                'g1-descriptive-stats': {'period_col': '日期'},
            },
        )

    def test_financial_model_report_counts_source_columns_not_runtime_aliases(self):
        import pandas as pd

        from services.finmod_compute import _build_analysis_data_stats

        result = _build_analysis_data_stats(
            pd.DataFrame(
                {
                    '编号': [1, 2],
                    '金额': [100, 200],
                    'x_i': [100, 200],
                }
            ),
            source_column_count=2,
            files=['数据.xlsx'],
            period_col='日期',
        )

        self.assertEqual(result['rows'], 2)
        self.assertEqual(result['columns'], 2)
        self.assertEqual(result['period_col'], '日期')
