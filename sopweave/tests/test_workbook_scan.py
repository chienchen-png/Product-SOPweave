from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from tools.workbook_scan import snapshot_and_scan, scan_workbook, read_workbook_values


class WorkbookScanTests(unittest.TestCase):
    def test_self_reference_uses_sheet_qualified_range_tokens_not_text_matches(self):
        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            book.active.title = "数据"
            book.active["A2"] = 10
            sheet = book.create_sheet("引用")
            for coordinate, formula in {
                "A2": "=数据!A2",
                "B2": '="B2"',
                "C2": "=SUM(C1:C3)",
                "D2": "='引用'!D2",
                "E2": "=[other.xlsx]引用!E2",
            }.items():
                sheet[coordinate] = formula
            path = Path(directory) / "references.xlsx"
            book.save(path)
            book.close()
            scanned = scan_workbook(path)["sheets"][1]
            self.assertEqual(
                {
                    item["cell"]: item["direct_self_reference"]
                    for item in scanned["formula_examples"]
                },
                {"A2": False, "B2": False, "C2": True, "D2": True, "E2": False},
            )

    def test_encrypted_ooxml_is_identified_before_the_zip_parser(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "encrypted.xlsx"
            path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"fixture")
            with patch("xlrd.compdoc.CompDoc") as compound:
                compound.return_value.get_named_stream.return_value = (
                    b"encryption metadata"
                )
                with self.assertRaisesRegex(ValueError, "加密"):
                    scan_workbook(path)

    def test_literal_equals_text_is_not_a_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.append(["code", "amount"])
            sheet.append(["=literal_text", 10])
            sheet["A2"].data_type = "s"
            path = Path(directory) / "literal.xlsx"
            book.save(path)
            book.close()
            manifest = scan_workbook(path)
            self.assertEqual(manifest["formula_count"], 0)
            self.assertFalse(manifest["has_unresolved_formulas"])
            self.assertEqual(read_workbook_values(path)["Sheet"][1][0], "=literal_text")

    def test_value_reader_ignores_formatting_tail_without_shifting_physical_coordinates(
        self,
    ):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = openpyxl.Workbook()
            sheet = book.active
            sheet["B3"], sheet["C3"] = "code", "amount"
            sheet["B4"], sheet["C4"] = "001", 0
            sheet["Z1000"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
            path = root / "styled.xlsx"
            book.save(path)
            book.close()
            rows = read_workbook_values(path)["Sheet"]
            self.assertEqual(len(rows), 4)
            self.assertEqual([len(row) for row in rows], [3, 3, 3, 3])
            self.assertEqual(rows[3], [None, "001", 0])

    def test_value_reader_rejects_workbook_cell_budget_and_row_limit_before_truncation(
        self,
    ):
        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            for sheet in (book.active, book.create_sheet("second")):
                sheet.append(["code", "amount"])
                sheet.append(["001", 1])
            path = Path(directory) / "budget.xlsx"
            book.save(path)
            book.close()
            with patch("tools.workbook_scan.MAX_IMPORT_CELLS", 7, create=True):
                with self.assertRaisesRegex(ValueError, "单元格"):
                    read_workbook_values(path)
            with self.assertRaisesRegex(ValueError, "行"):
                read_workbook_values(path, max_rows=1)

    def _make_workbook(self, root: Path) -> Path:
        path = root / "scan.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "数据"
        ws.append(["编号", "金额"])
        ws.append(["001", 3])
        ws["B3"] = "=B2*2"
        ws["Z100"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        ws.row_dimensions[2].hidden = True
        ws.column_dimensions["B"].hidden = True
        table = Table(displayName="DataTable", ref="A1:B3")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

        notes = wb.create_sheet("说明")
        notes.merge_cells("A1:C1")
        notes["A1"] = "使用说明"
        wb.save(path)
        wb.close()
        return path

    def test_scan_records_formula_structure_and_actual_bounds(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._make_workbook(Path(tmp))
            manifest = scan_workbook(path)

            from tools.workbook_scan import MANIFEST_SCHEMA_VERSION

            self.assertEqual(manifest["schema_version"], MANIFEST_SCHEMA_VERSION)
        self.assertEqual(manifest["format"], "xlsx")
        self.assertEqual(manifest["formula_count"], 1)
        self.assertEqual(manifest["formula_missing_count"], 1)
        self.assertTrue(manifest["has_unresolved_formulas"])

        data = manifest["sheets"][0]
        self.assertEqual(data["name"], "数据")
        self.assertEqual(data["declared_range"], "A1:Z100")
        self.assertEqual(data["actual_range"], "A1:B3")
        self.assertEqual(data["formula_count"], 1)
        self.assertEqual(data["formula_missing_count"], 1)
        self.assertEqual(data["formula_examples"][0]["cell"], "B3")
        self.assertEqual(data["formula_examples"][0]["formula"], "=B2*2")
        self.assertEqual(data["hidden_rows"], [2])
        self.assertEqual(data["hidden_columns"], ["B"])
        self.assertEqual(data["tables"], [{"name": "DataTable", "ref": "A1:B3"}])

        notes = manifest["sheets"][1]
        self.assertEqual(notes["actual_range"], "A1")
        self.assertEqual(notes["merged_ranges"], ["A1:C1"])

    def test_snapshot_is_content_addressed_and_manifest_is_portable(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = self._make_workbook(root)
            first = snapshot_and_scan(path, root)
            second = snapshot_and_scan(path, root)

            self.assertEqual(first["file_hash"], second["file_hash"])
            cache_dir = root / ".sopweave" / "import-cache" / first["file_hash"]
            self.assertEqual(Path(first["snapshot_path"]), (cache_dir / "original.xlsx").resolve())
            self.assertTrue((cache_dir / "original.xlsx").is_file())
            self.assertTrue((cache_dir / "manifest.json").is_file())

            persisted = json.loads(
                (cache_dir / "manifest.json").read_text(encoding="utf-8")
            )
            serialized = json.dumps(persisted, ensure_ascii=False)
            self.assertNotIn(str(root.resolve()), serialized)
            self.assertEqual(persisted["source_name"], "scan.xlsx")

    def test_importer_blocks_unverified_formula_values(self) -> None:
        from tools.project_data_importer import (
            FormulaCalculationRequired,
            ProjectDataImporter,
        )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "formula.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据"
            ws.append(["编号", "金额"])
            ws.append(["001", "=1+2"])
            wb.save(path)
            wb.close()

            db = sqlite3.connect(":memory:")
            db.row_factory = sqlite3.Row
            schema = (
                Path(__file__).resolve().parents[1] / "storage" / "schema.sql"
            ).read_text(encoding="utf-8")
            db.executescript(schema)
            blocked = {
                "requires_review": True,
                "missing_after": 1,
                "missing_before": 1,
                "attempts": [],
            }
            with patch(
                "tools.project_data_importer.ensure_calculated_workbook",
                return_value=blocked,
            ):
                with self.assertRaises(FormulaCalculationRequired) as caught:
                    ProjectDataImporter(db).import_workbook(
                        "project", path.name, str(root), anchor_rows={"数据": 0}
                    )
            self.assertEqual(caught.exception.calculation["missing_after"], 1)
            self.assertEqual(
                db.execute("SELECT COUNT(*) FROM project_data_files").fetchone()[0],
                0,
            )
            snapshots = list(
                (root / ".sopweave" / "import-cache").glob("*/original.xlsx")
            )
            self.assertEqual(len(snapshots), 1)
            db.close()

    def test_legacy_import_cache_directory_migrates_once(self) -> None:
        from tools.workbook_scan import snapshot_and_scan

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            legacy = root / ".finance-agent" / "import-cache"
            legacy.mkdir(parents=True)
            (legacy / "sentinel.txt").write_text("legacy", encoding="utf-8")
            path = self._make_workbook(root)

            result = snapshot_and_scan(path, root)

            self.assertFalse((root / ".finance-agent").exists())
            self.assertTrue((root / ".sopweave" / "import-cache" / "sentinel.txt").is_file())
            self.assertTrue(
                (root / ".sopweave" / "import-cache" / result["file_hash"] / "original.xlsx").is_file()
            )

    def test_importer_records_verified_cached_value_pipeline(self) -> None:
        from tools.project_data_importer import ProjectDataImporter

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "values.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据"
            ws.append(["编号", "金额"])
            ws.append(["001", 3])
            wb.save(path)
            wb.close()

            db = sqlite3.connect(":memory:")
            db.row_factory = sqlite3.Row
            schema = (
                Path(__file__).resolve().parents[1] / "storage" / "schema.sql"
            ).read_text(encoding="utf-8")
            db.executescript(schema)
            result = ProjectDataImporter(db).import_workbook(
                "project", path.name, str(root), anchor_rows={"数据": 0}
            )
            row = db.execute(
                "SELECT meta_json FROM project_data_files WHERE id = ?",
                (result["file_id"],),
            ).fetchone()
            meta = json.loads(row["meta_json"])
            from tools.project_data_importer import IMPORT_VERSION

            self.assertEqual(meta["import_version"], IMPORT_VERSION)
            self.assertTrue(meta["formulas_materialized"])
            self.assertEqual(meta["formula_audit"]["engine"], "cached-values")
            self.assertEqual(meta["formula_audit"]["missing_count"], 0)
            db.close()

    def test_scan_detects_direct_self_reference(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "circular.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = "值"
            ws["A2"] = "=A2+1"
            wb.save(path)
            wb.close()

            manifest = scan_workbook(path)

        sheet = manifest["sheets"][0]
        self.assertEqual(sheet["direct_self_reference_count"], 1)
        self.assertTrue(sheet["formula_examples"][0]["direct_self_reference"])


if __name__ == "__main__":
    unittest.main()
