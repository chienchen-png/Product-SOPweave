from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

import openpyxl

from tools.project_data_importer import ProjectDataImporter, _apply_layout_override
from tools.project_data_tools import _export_precise
from tools.table_layout_detector import detect_layout
from storage.project_data_store import ProjectDataStore


class ImportLayoutTests(unittest.TestCase):
    def test_single_column_and_sparse_trailing_data_are_never_discarded_as_notes(self):
        for rows in [[["金额"], [100]], [["编号"], ["00123"], ["00001"]]]:
            with self.subTest(rows=rows):
                layout = detect_layout(rows)
                self.assertEqual(layout.data_end_row, len(rows) - 1)
                self.assertEqual(layout.classification, "table")
        sparse = detect_layout([["编号", "金额"], ["A", 10], ["B", None]])
        self.assertEqual(sparse.data_end_row, 2)
        self.assertFalse(sparse.is_confident)

    def test_numeric_parent_headers_preserve_vertical_key_and_both_year_groups(self):
        rows = [
            ["项目", "2025", None, "2026", None],
            [None, "收入", "成本", "收入", "成本"],
            ["P1", 100, 60, 120, 70],
        ]
        layout = detect_layout(rows, {"merged_ranges": ["A1:A2", "B1:C1", "D1:E1"]})
        self.assertEqual(layout.header_row, 0)
        self.assertEqual(
            layout.columns, ["项目", "2025_收入", "2025_成本", "2026_收入", "2026_成本"]
        )
        self.assertEqual(
            (layout.start_col, layout.end_col, layout.data_start_row), (0, 4, 2)
        )

    def test_separated_regions_with_different_headers_require_review(self):
        rows = [
            ["编号", "金额"],
            ["A", 10],
            [None, None],
            ["其他区", "金额"],
            ["B", 99],
        ]
        layout = detect_layout(rows)
        self.assertFalse(layout.is_confident)
        self.assertEqual(layout.data_end_row, 1)
        self.assertEqual(layout.region_source, "multiple_regions")
        continuous = detect_layout(
            [["编号", "金额"], ["A", 10], [None, None], ["B", 99]]
        )
        self.assertEqual(continuous.data_end_row, 3)
        self.assertNotEqual(continuous.region_source, "multiple_regions")

    def test_confirming_the_detected_merged_header_keeps_hierarchical_labels(self):
        rows = [["地区", "收入", None], [None, "金额", "税额"], ["华东", 100, 10]]
        layout = detect_layout(rows, {"merged_ranges": ["B1:C1"]})
        _apply_layout_override(
            layout,
            rows,
            {"header_row": 0, "data_start_row": 2, "start_col": 0, "end_col": 2},
        )
        self.assertEqual(layout.columns, ["地区", "收入_金额", "收入_税额"])
        self.assertEqual(layout.data_start_row, 2)

    def test_multiple_regions_require_review_and_suggest_only_the_first_region(self):
        rows = [["编号", "金额"], ["A", 10], [None, None], ["编号", "金额"], ["B", 20]]
        layout = detect_layout(rows)
        self.assertFalse(layout.is_confident)
        self.assertEqual(layout.data_end_row, 1)
        self.assertEqual(layout.region_source, "multiple_regions")

    def test_multiple_native_tables_require_explicit_region_review(self):
        rows = [
            ["code", "amount"],
            ["A", 10],
            [None, None],
            ["code", "amount"],
            ["B", 20],
        ]
        layout = detect_layout(
            rows,
            {
                "tables": [
                    {"name": "first", "ref": "A1:B2"},
                    {"name": "second", "ref": "A4:B5"},
                ]
            },
        )
        self.assertFalse(layout.is_confident)
        self.assertEqual(layout.region_source, "multiple_excel_tables")

    def test_manual_region_rejects_fractional_out_of_bounds_and_empty_selections(self):
        rows = [["code", "amount"], ["001", 10], ["002", 20]]
        invalid = [
            {"header_row": 0.5},
            {"header_row": True},
            {"data_start_row": 99},
            {"data_end_row": 99},
            {"data_end_row": -2},
            {"start_col": 99},
            {"end_col": 99},
            {"excluded_columns": [-1]},
            {"excluded_columns": [2]},
            {"excluded_columns": [0, 1]},
            {"excluded_columns": [0.5]},
        ]
        for override in invalid:
            with self.subTest(override=override):
                layout = detect_layout(rows)
                with self.assertRaises(ValueError):
                    _apply_layout_override(layout, rows, {"header_row": 0, **override})

    def _db(self) -> sqlite3.Connection:
        db = sqlite3.connect(":memory:")
        db.row_factory = sqlite3.Row
        schema = (
            Path(__file__).resolve().parents[1] / "storage" / "schema.sql"
        ).read_text(encoding="utf-8")
        db.executescript(schema)
        return db

    def test_trailing_unnamed_note_column_is_excluded_and_rows_keep_lineage(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "sparse.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据"
            ws.append(["编号", "金额", None])
            ws.append(["001", 10, "零散说明"])
            ws.append([None, None, None])
            ws.append(["002", 20, None])
            wb.save(path)
            wb.close()

            db = self._db()
            result = ProjectDataImporter(db).import_workbook(
                "project", path.name, str(root), anchor_rows={"数据": 0}
            )
            sheet = db.execute(
                "SELECT * FROM project_data_sheets WHERE file_id = ?",
                (result["file_id"],),
            ).fetchone()
            columns = [
                row["name"]
                for row in db.execute(
                    f'PRAGMA table_info("{sheet["table_name"]}")'
                ).fetchall()
            ]
            rows = db.execute(
                f'SELECT _source_row, "编号", "金额" '
                f'FROM "{sheet["table_name"]}" ORDER BY id'
            ).fetchall()
            meta = json.loads(
                db.execute(
                    "SELECT meta_json FROM project_data_files WHERE id = ?",
                    (result["file_id"],),
                ).fetchone()["meta_json"]
            )

            self.assertEqual(
                columns,
                ["id", "_source_row", "_source_region", "编号", "金额"],
            )
            self.assertEqual([row["_source_row"] for row in rows], [2, 4])
            self.assertEqual(meta["layouts"]["数据"]["excluded_columns"], [2])
            headers = json.loads(sheet["headers_json"])
            self.assertEqual(
                [(item["excel_col"], item["source_col"]) for item in headers],
                [("A", 0), ("B", 1)],
            )
            db.close()

    def test_merged_group_header_uses_real_merge_structure(self) -> None:
        rows = [
            ["地区", "收入", None],
            [None, "金额", "税额"],
            ["华东", 100, 10],
        ]
        layout = detect_layout(
            rows,
            {"merged_ranges": ["B1:C1"], "tables": []},
        )
        self.assertEqual(layout.columns, ["地区", "收入_金额", "收入_税额"])
        self.assertEqual(layout.data_start_row, 2)

    def test_named_range_is_preferred_over_surrounding_notes(self) -> None:
        rows = [
            ["月度收入说明", None, None, None],
            [None, "编号", "金额", None],
            [None, "001", 100, "范围外备注"],
            [None, "002", 200, None],
        ]
        layout = detect_layout(
            rows,
            {
                "tables": [],
                "named_ranges": [{"name": "收入数据", "ref": "B2:C4"}],
            },
            "收入",
        )
        self.assertEqual(layout.columns, ["编号", "金额"])
        self.assertEqual((layout.start_col, layout.end_col), (1, 2))
        self.assertEqual(layout.region_source, "named_range")

    def test_instruction_sheet_is_classified_as_notes(self) -> None:
        rows = [["使用说明"], ["请填写黄色单元格"]]
        layout = detect_layout(rows, {"tables": []}, "说明")
        self.assertEqual(layout.classification, "notes")

    def test_precise_writeback_uses_physical_row_and_column(self) -> None:
        class Writer:
            def __init__(self) -> None:
                self.cells: list[tuple[str, int, int, object]] = []

            def has(self, _sheet: str) -> bool:
                return True

            def set_cell(self, sheet: str, row: int, col: int, value: object) -> None:
                self.cells.append((sheet, row, col, value))

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "offset.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据"
            ws.append([None, "编号", "金额", None])
            ws.append([None, "001", 10, "边界外说明"])
            ws.append([None, None, None, None])
            ws.append([None, "002", 20, None])
            wb.save(path)
            wb.close()

            db = self._db()
            result = ProjectDataImporter(db).import_workbook(
                "project", path.name, str(root), anchor_rows={"数据": 0}
            )
            store = ProjectDataStore(db)
            sheet = store.list_sheets(result["file_id"])[0]
            second = db.execute(
                f'SELECT id FROM "{sheet["table_name"]}" ORDER BY id LIMIT 1 OFFSET 1'
            ).fetchone()
            store.log_change(
                result["file_id"],
                "数据",
                second["id"],
                "update",
                col_index=1,
                old_value=20,
                new_value=25,
                col_name="金额",
            )
            writer = Writer()
            written = _export_precise(writer, store, store.get_file(result["file_id"]))
            self.assertEqual(written, 1)
            self.assertEqual(writer.cells, [("数据", 4, 3, 25)])
            db.close()

    def test_user_region_override_excludes_selected_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "region.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "数据"
            ws.append(["编号", "内部备注", "金额"])
            ws.append(["001", "忽略", 10])
            ws.append(["002", "忽略", 20])
            wb.save(path)
            wb.close()
            db = self._db()
            result = ProjectDataImporter(db).import_workbook(
                "project",
                path.name,
                str(root),
                anchor_rows={
                    "数据": {
                        "header_row": 0,
                        "data_start_row": 1,
                        "data_end_row": 2,
                        "start_col": 0,
                        "end_col": 2,
                        "excluded_columns": [1],
                    },
                },
            )
            sheet = db.execute(
                "SELECT * FROM project_data_sheets WHERE file_id = ?",
                (result["file_id"],),
            ).fetchone()
            headers = json.loads(sheet["headers_json"])
            self.assertEqual([item["name"] for item in headers], ["编号", "金额"])
            self.assertEqual([item["excel_col"] for item in headers], ["A", "C"])
            rows = db.execute(
                f'SELECT "编号", "金额" FROM "{sheet["table_name"]}" ORDER BY id'
            ).fetchall()
            self.assertEqual([tuple(row) for row in rows], [("001", 10), ("002", 20)])
            db.close()


if __name__ == "__main__":
    unittest.main()
