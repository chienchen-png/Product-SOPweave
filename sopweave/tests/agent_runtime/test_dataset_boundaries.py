from __future__ import annotations

import asyncio
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

import openpyxl

from agent_runtime.context import WorkbenchRunContext
from agent_runtime.definitions.agent_registry import AgentDefinition, AgentManifest
from agent_runtime.tools.catalog import ToolCatalog
from tests.agent_runtime.test_tool_catalog import _invoke
from tools.project_data_importer import ProjectDataImporter


class DatasetBoundaryTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.path = self.root/'db.sqlite'
        with closing(self.connect()) as db:
            db.executescript(Path('storage/schema.sql').read_text(encoding='utf8'))
        self.context = WorkbenchRunContext(
            project_id='own', work_dir=self.root, run_id='boundary-run',
            conversation_id='boundary-chat', work_mode='auto',
            agent_id='general-agent', database_factory=self.connect,
        )
        definition = AgentDefinition(AgentManifest(
            id='general-agent', name='test', version='1', toolsets=['project-data'],
        ), '', self.root)
        self.tools = {tool.name:tool for tool in ToolCatalog.default().build(definition,self.context)}

    def connect(self):
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        return db

    def workbook(self, name, rows):
        book = openpyxl.Workbook()
        for row in rows:
            book.active.append(row)
        book.save(self.root/name)
        book.close()

    def invoke(self, name, **arguments):
        return asyncio.run(_invoke(self.tools[name],self.context,arguments))

    def test_scoped_catalog_composite_filters_aggregation_and_transaction_boundaries(self):
        self.workbook('own.xlsx',[
            ['code','dept','amount'], ['A','X',10], ['A','Y',20], ['B','X',30],
        ])
        self.workbook('other.xlsx',[['code','amount'],['PRIVATE',999]])
        with closing(self.connect()) as db:
            importer = ProjectDataImporter(db)
            importer.import_workbook('own','own.xlsx',str(self.root),anchor_rows={'Sheet':0})
            importer.import_workbook('other','other.xlsx',str(self.root),anchor_rows={'Sheet':0})
        catalog = self.invoke('list_datasets')
        self.assertEqual([f['file'] for f in catalog['files']],['own.xlsx'])
        self.assertNotIn('PRIVATE',str(catalog))
        self.assertEqual(catalog['files'][0]['sheets'][0]['rows'],3)
        self.assertNotIn('fields',catalog['files'][0]['sheets'][0])
        fields = self.invoke('list_datasets',file='own.xlsx',sheet='Sheet')
        self.assertEqual(
            [item['name'] for item in fields['files'][0]['sheets'][0]['fields']],
            ['code','dept','amount'],
        )
        filters = [{'field':'code','value':'A'},{'field':'dept','value':'Y'}]
        found = self.invoke('query_dataset',file='own.xlsx',columns=['code','dept','amount'],filters=filters,limit=1)
        self.assertEqual(found['rows'],[{'code':'A','dept':'Y','amount':20}])
        self.assertTrue(found['revision'])
        limited = self.invoke('query_dataset',file='own.xlsx',limit=1)
        self.assertEqual(len(limited['rows']),1)
        self.assertTrue(limited['truncated'])
        total = self.invoke('query_dataset',file='own.xlsx',metrics=[{'function':'sum','field':'amount'}])
        self.assertEqual(total['rows'],[{'sum(amount)':60}])
        denied = self.invoke('apply_data_changes',file='own.xlsx',operation='delete',expected_revision=found['revision'])
        self.assertFalse(denied['success'])
        unchanged = self.invoke('query_dataset',file='own.xlsx')
        self.assertEqual(unchanged['revision'],found['revision'])
        changed = self.invoke('apply_data_changes',file='own.xlsx',operation='update',
            expected_revision=found['revision'],filters=filters,values=[{'field':'amount','value':25}])
        self.assertTrue(changed['success'],changed)
        after = self.invoke('query_dataset',file='own.xlsx')
        self.assertEqual([r['amount'] for r in after['rows']],[10,25,30])
        with closing(self.connect()) as db:
            audit = db.execute('select old_value,new_value,run_id from project_data_changes').fetchall()
        self.assertEqual(len(audit),1)
        self.assertEqual(audit[0]['run_id'],'boundary-run')
        self.assertEqual((str(audit[0]['old_value']),str(audit[0]['new_value'])),('20','25'))

    def test_ambiguous_import_returns_review_request_without_persisting_partial_data(self):
        self.workbook('review.xlsx',[
            ['编号','金额'],['A',10],[None,None],['其他区','金额'],['B',99],
        ])
        book = openpyxl.load_workbook(self.root/'review.xlsx')
        from openpyxl.worksheet.table import Table
        book.active.add_table(Table(displayName='First',ref='A1:B2'))
        book.active.add_table(Table(displayName='Second',ref='A4:B5'))
        book.save(self.root/'native.xlsx')
        book.close()
        for filename in ('review.xlsx','native.xlsx'):
            with self.subTest(filename=filename):
                result = self.invoke('import_dataset',path=filename)
                self.assertFalse(result['success'])
                self.assertTrue(result.get('needs_layout_review'),result)
                self.assertIn('数据导入',result['message'])
                confirmed = self.invoke('import_dataset',path=filename,header_row=1)
                self.assertFalse(confirmed['success'])
                self.assertTrue(confirmed.get('needs_layout_review'),confirmed)
        self.assertEqual(self.invoke('list_datasets')['files'],[])

    def test_catalog_stays_compact_until_one_file_is_requested(self):
        rows = [['code'] + [f'field_{index}' for index in range(20)]]
        rows.extend([['A'] + list(range(20))])
        book = openpyxl.Workbook()
        book.active.title = 'Sheet_0'
        for row in rows:
            book.active.append(row)
        for index in range(1, 12):
            sheet = book.create_sheet(f'Sheet_{index}')
            sheet.append(rows[0])
            sheet.append(rows[1])
        book.save(self.root/'catalog.xlsx')
        book.close()
        with closing(self.connect()) as db:
            ProjectDataImporter(db).import_workbook(
                'own',
                'catalog.xlsx',
                str(self.root),
                anchor_rows={f'Sheet_{index}': 0 for index in range(12)},
            )

        catalog = self.invoke('list_datasets')
        target = next(item for item in catalog['files'] if item['file'] == 'catalog.xlsx')
        self.assertEqual(target['sheet_count'], 12)
        self.assertTrue(all('fields' not in item for item in target['sheets']))
        self.assertLess(len(str(catalog).encode('utf-8')), 10000)

        detail = self.invoke(
            'list_datasets',
            file='catalog.xlsx',
            sheet='Sheet_7',
        )
        sheet = detail['files'][0]['sheets'][0]
        self.assertEqual(sheet['sheet'], 'Sheet_7')
        self.assertEqual(len(sheet['fields']), 21)

    def test_user_header_row_is_one_based_and_preserves_first_data_row(self):
        self.workbook('sparse.xlsx',[['code','amount'],['A',10],['B',None]])
        pending = self.invoke('import_dataset',path='sparse.xlsx')
        self.assertTrue(pending['needs_anchor'],pending)
        confirmed = self.invoke('import_dataset',path='sparse.xlsx',header_row=1)
        self.assertTrue(confirmed['success'],confirmed)
        result = self.invoke('query_dataset',file='sparse.xlsx',columns=['code','amount'])
        self.assertTrue(result['success'],result)
        self.assertEqual(result['rows'],[{'code':'A','amount':10},{'code':'B','amount':None}])
