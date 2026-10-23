import unittest
import sqlite3
import tempfile
from contextlib import closing
from pathlib import Path
from decimal import Decimal
from services.money import parse_amount
from services.skills.policies import compile_policy, SkillOperation
from services.skills.contracts import SkillProblem
from tests.agent_runtime.test_skill_contracts import sample_spec


class SkillPolicyTests(unittest.TestCase):
    def policy(self, **kwargs):
        return compile_policy(
            sample_spec(
                required_tools=["read_sheet_range", "write_text", "edit_workbook"],
                execution_policy=kwargs,
            )
        )

    def test_decimal_tolerance_and_rounding_are_explicit(self):
        self.assertEqual(
            parse_amount("270.00") - parse_amount("170.00"), Decimal("100.00")
        )
        self.assertEqual(parse_amount("0.30") - parse_amount("0.29"), Decimal("0.01"))
        for value in ("NaN", "Infinity", "bad"):
            with self.assertRaises(ValueError):
                parse_amount(value)

    def test_auto_mode_cannot_skip_required_approval(self):
        policy = self.policy(
            require_approval_for=["edit_workbook"], write_paths=["out.xlsx"]
        )
        decision = policy.evaluate(
            SkillOperation("edit_workbook", "write", {"path": "out.xlsx"}, Path.cwd())
        )
        self.assertEqual(decision.action, "confirm")

    def test_amount_precision_is_enforced_before_office_write(self):
        policy = self.policy(write_paths=['out.xlsx'],amount_fields=['Sheet!A1'],decimal_places=2)
        for value,action in [('12.34','allow'),('12.345','deny'),('NaN','deny')]:
            with self.subTest(value=value):
                decision = policy.evaluate(SkillOperation('edit_workbook','write',{
                    'path':'out.xlsx','changes':[{'sheet':'Sheet','cell':'A1','value':value}],
                },Path.cwd()))
                self.assertEqual(decision.action,action)

    def test_unrestricted_command_cannot_bypass_controlled_sop(self):
        self.assertEqual(
            self.policy()
            .evaluate(SkillOperation("run_process", "command", {}, Path.cwd()))
            .action,
            "deny",
        )

    def test_write_text_cannot_overwrite_protected_source(self):
        self.assertEqual(
            self.policy(write_paths=["report.md"])
            .evaluate(
                SkillOperation(
                    "write_text", "write", {"path": "source.xlsx"}, Path.cwd()
                )
            )
            .action,
            "deny",
        )

    def test_trailing_slash_write_directory_allows_only_its_descendants(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "reports").mkdir()
            policy = self.policy(write_paths=["reports/"])
            self.assertEqual(
                policy.evaluate(
                    SkillOperation(
                        "write_text",
                        "write",
                        {"path": "reports/month-end.md"},
                        root,
                    )
                ).action,
                "allow",
            )
            for path in ("reports-evil/month-end.md", "month-end.md"):
                with self.subTest(path=path):
                    self.assertEqual(
                        policy.evaluate(
                            SkillOperation(
                                "write_text", "write", {"path": path}, root
                            )
                        ).action,
                        "deny",
                    )

    def test_unknown_critical_rule_blocks_publish(self):
        with self.assertRaises(SkillProblem):
            self.policy(checks=["invented_tax_validator"])

    def test_missing_required_input_blocks_business_write(self):
        decision = self.policy(
            required_inputs=["currency"], write_paths=["report.md"]
        ).evaluate(
            SkillOperation("write_text", "write", {"path": "report.md"}, Path.cwd())
        )
        self.assertEqual(decision.action, "deny")

    def test_stale_source_version_blocks_writeback(self):
        from agent_runtime.workspace import file_version

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.txt"
            path.write_text("first")
            policy = self.policy(
                write_paths=["report.md"],
                source_versions=[dict(path="source.txt", version=file_version(path))],
            )
            path.write_text("changed")
            result = policy.evaluate(
                SkillOperation(
                    "write_text", "write", {"path": "report.md"}, Path(folder)
                )
            )
            self.assertEqual(result.action, "deny")
            self.assertIn("版本冲突", result.reason)

    def test_real_reconcile_decimal_boundary_and_zero_tolerance(self):
        import openpyxl
        from tools.project_data_importer import ProjectDataImporter
        from tools.reconcile_tool import reconcile_variables
        from services.reconciliation import _resolve_verify_params
        from storage.project_data_store import ProjectDataStore

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with closing(sqlite3.connect(root / "test.db")) as db:
                db.row_factory = sqlite3.Row
                db.executescript(Path("storage/schema.sql").read_text(encoding="utf-8"))
                for name, amount in [("base.xlsx", "0.30"), ("detail.xlsx", "0.29")]:
                    book = openpyxl.Workbook()
                    book.active.append(["code", "amount"])
                    book.active.append(["A", amount])
                    book.save(root / name)
                    book.close()
                    ProjectDataImporter(db).import_workbook(
                        "p", name, str(root), anchor_rows={"Sheet": 0}
                    )
                result = reconcile_variables(
                    "base.xlsx",
                    "detail.xlsx",
                    match_keys=["code"],
                    amount_col="amount",
                    tolerance="0.01",
                    project_id="p",
                    db_override=db,
                )
                self.assertTrue(result["success"], result)
                self.assertEqual(result["summary"]["ok"], 1)
                zero = reconcile_variables(
                    "base.xlsx",
                    "detail.xlsx",
                    match_keys=["code"],
                    amount_col="amount",
                    tolerance=0,
                    project_id="p",
                    db_override=db,
                )
                self.assertEqual(zero["summary"]["b"], 1)
                self.assertEqual(zero["differences"][0]["diff"], 0.01)
                params = dict(
                    matchKeys={"k1": "code"},
                    variables=["amount"],
                    baseFile="base.xlsx",
                    detailFile="detail.xlsx",
                    baseAmountCol="amount",
                    tolerance=0,
                )
                self.assertEqual(
                    _resolve_verify_params(params, ProjectDataStore(db), "p")[
                        "tolerance"
                    ],
                    0,
                )

    def test_reconcile_does_not_silently_drop_identity_conflicts(self):
        import openpyxl
        from tools.project_data_importer import ProjectDataImporter
        from tools.reconcile_tool import reconcile_variables

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with closing(sqlite3.connect(root / "test.db")) as db:
                db.row_factory = sqlite3.Row
                db.executescript(
                    Path("storage/schema.sql").read_text(encoding="utf-8")
                )
                for name, rows in [
                    ("base.xlsx", [["K1", "A", "100"]]),
                    (
                        "detail.xlsx",
                        [
                            ["K1", "B", "100"],
                            ["K1", "C", "100"],
                        ],
                    ),
                ]:
                    book = openpyxl.Workbook()
                    book.active.append(["contract", "person", "amount"])
                    for row in rows:
                        book.active.append(row)
                    book.save(root / name)
                    book.close()
                    ProjectDataImporter(db).import_workbook(
                        "p", name, str(root), anchor_rows={"Sheet": 0}
                    )

                result = reconcile_variables(
                    "base.xlsx",
                    "detail.xlsx",
                    match_keys=["contract", "person"],
                    amount_col="amount",
                    project_id="p",
                    db_override=db,
                )

                self.assertTrue(result["success"], result)
                self.assertEqual(result["summary"]["total"], 1)
                self.assertEqual(result["summary"]["b"], 1)
                self.assertEqual(result["summary"]["matched"], 0)
                self.assertEqual(result["differences"][0]["key1"], "K1")
                self.assertIn("身份", result["differences"][0]["reason"])

    def test_reconcile_rejects_negative_tolerance(self):
        import openpyxl
        from tools.project_data_importer import ProjectDataImporter
        from services.reconciliation import _resolve_verify_params
        from storage.project_data_store import ProjectDataStore

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with closing(sqlite3.connect(root / "test.db")) as db:
                db.row_factory = sqlite3.Row
                db.executescript(
                    Path("storage/schema.sql").read_text(encoding="utf-8")
                )
                for name in ("base.xlsx", "detail.xlsx"):
                    book = openpyxl.Workbook()
                    book.active.append(["contract", "amount"])
                    book.active.append(["K1", "100"])
                    book.save(root / name)
                    book.close()
                    ProjectDataImporter(db).import_workbook(
                        "p", name, str(root), anchor_rows={"Sheet": 0}
                    )
                params = {
                    "matchKeys": {"k1": "contract"},
                    "variables": ["amount"],
                    "baseFile": "base.xlsx",
                    "detailFile": "detail.xlsx",
                    "baseAmountCol": "amount",
                    "tolerance": -1,
                }

                with self.assertRaisesRegex(ValueError, "容差"):
                    _resolve_verify_params(params, ProjectDataStore(db), "p")

    def test_reconcile_three_key_levels_and_person_customer_fallback(self):
        import openpyxl
        from tools.project_data_importer import ProjectDataImporter
        from tools.reconcile_tool import reconcile_variables

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with closing(sqlite3.connect(root / "test.db")) as db:
                db.row_factory = sqlite3.Row
                db.executescript(
                    Path("storage/schema.sql").read_text(encoding="utf-8")
                )
                fixtures = {
                    "base3.xlsx": [
                        ["C1", "A", "X", "100"],
                        ["C1", "A", "Y", "200"],
                        ["C1", "B", "Z", "300"],
                    ],
                    "detail3.xlsx": [
                        ["C1", "A", "X", "100"],
                        ["C1", "A", "Y", "200"],
                        ["C1", "B", "Z", "300"],
                    ],
                    "base-fallback.xlsx": [["K9", "A", "X", "400"]],
                    "detail-fallback.xlsx": [["K8", "A", "X", "400"]],
                }
                for name, rows in fixtures.items():
                    book = openpyxl.Workbook()
                    book.active.append(["contract", "person", "customer", "amount"])
                    for row in rows:
                        book.active.append(row)
                    book.save(root / name)
                    book.close()
                    ProjectDataImporter(db).import_workbook(
                        "p", name, str(root), anchor_rows={"Sheet": 0}
                    )

                cascade = reconcile_variables(
                    "base3.xlsx",
                    "detail3.xlsx",
                    match_keys=["contract", "person", "customer"],
                    amount_col="amount",
                    project_id="p",
                    db_override=db,
                )
                self.assertTrue(cascade["success"], cascade)
                self.assertEqual(cascade["summary"]["ok"], 3)

                fallback = reconcile_variables(
                    "base-fallback.xlsx",
                    "detail-fallback.xlsx",
                    match_keys=["contract", "person", "customer"],
                    amount_col="amount",
                    unmatched_action="fallback_person_customer",
                    project_id="p",
                    db_override=db,
                )
                self.assertTrue(fallback["success"], fallback)
                self.assertEqual(fallback["summary"]["ok"], 1)
                self.assertEqual(
                    fallback["fallback_stats"], {"attempted": 1, "found": 1}
                )
