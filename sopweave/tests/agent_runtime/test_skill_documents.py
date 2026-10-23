from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from services.skills.documents import parse_source, ParseLimits
from services.skills.contracts import SkillProblem


class SkillDocumentTests(unittest.TestCase):
    def test_docx_tables_keep_step_order(self):
        with tempfile.TemporaryDirectory(prefix="文档 ") as folder:
            path = Path(folder) / "核对流程.docx"
            document = Document()
            document.add_paragraph("第一步：读取流水")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "币种"
            table.cell(0, 1).text = "人民币"
            document.add_paragraph("第三步：生成报告")
            document.save(path)
            source = parse_source(path)
            self.assertLess(source.text.index("读取流水"), source.text.index("人民币"))
            self.assertLess(source.text.index("人民币"), source.text.index("生成报告"))
            self.assertTrue(source.complete)
            self.assertGreaterEqual(len(source.locations), 3)

    def test_pdf_keeps_page_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "text.pdf"
            writer = PdfWriter()
            page = writer.add_blank_page(width=200, height=200)
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 10 100 Td (SOP reconcile) Tj ET")
            page[NameObject("/Contents")] = stream
            with path.open("wb") as output:
                writer.write(output)
            source = parse_source(path)
            self.assertIn("SOP reconcile", source.text)
            self.assertEqual(source.locations[0]["page"], 1)

    def test_scanned_pdf_is_reported_unsupported(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with path.open("wb") as output:
                writer.write(output)
            with self.assertRaisesRegex(SkillProblem, "OCR|文本"):
                parse_source(path)

    def test_corrupt_encrypted_oversize_and_timeout_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.docx"
            path.write_bytes(b"bad")
            with self.assertRaises(SkillProblem):
                parse_source(path)
            path = Path(folder) / "big.txt"
            path.write_text("x" * 100, encoding="utf-8")
            with self.assertRaises(SkillProblem):
                parse_source(path, limits=ParseLimits(max_file_bytes=10))
            with self.assertRaises(SkillProblem):
                parse_source(path, limits=ParseLimits(timeout_seconds=0.001))

    def test_unicode_text_and_filename(self):
        with tempfile.TemporaryDirectory(prefix="中文 空格") as folder:
            path = Path(folder) / "财务规则.md"
            path.write_text("月末按照人民币进行核对", encoding="utf-8-sig")
            self.assertIn("人民币", parse_source(path).text)
