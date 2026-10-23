from __future__ import annotations

import asyncio
import sqlite3
import tempfile
import unittest
from pathlib import Path

import openpyxl

from agent_runtime.context import WorkbenchRunContext
from agent_runtime.definitions.agent_registry import AgentDefinition, AgentManifest
from agent_runtime.tools.catalog import ToolCatalog
from tests.agent_runtime.test_tool_catalog import _invoke
from tools.project_data_importer import ProjectDataImporter


class DatasetLifecycleTests(unittest.TestCase):
    def test_visible_and_hidden_projects_share_typed_modify_and_export(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        directory = temporary.name
        root = Path(directory)
        database = root / "test.db"

        def factory():
            db = sqlite3.connect(database)
            db.row_factory = sqlite3.Row
            return db

        db = factory()
        self.addCleanup(db.close)
        db.executescript(Path("storage/schema.sql").read_text(encoding="utf-8"))
        for index, project in enumerate(("visible", "__app_recon__", "__app_finmod__")):
            name = f"data{index}.xlsx"
            path = root / name
            book = openpyxl.Workbook()
            book.active.append(["code", "amount"])
            book.active.append(["A", 10])
            book.save(path)
            book.close()
            ProjectDataImporter(db).import_workbook(
                project, name, str(root), anchor_rows={"Sheet": 0}
            )
            context = WorkbenchRunContext(
                project_id=project,
                work_dir=root,
                run_id=f"r{index}",
                conversation_id="c",
                work_mode="auto",
                agent_id="general-agent",
                database_factory=factory,
            )
            definition = AgentDefinition(
                AgentManifest(
                    id="general-agent",
                    name="test",
                    version="1",
                    toolsets=["project-data"],
                ),
                "",
                root,
            )
            tools = {
                tool.name: tool
                for tool in ToolCatalog.default().build(definition, context)
            }

            async def flow():
                result = await _invoke(tools["query_dataset"], context, {"file": name})
                changed = await _invoke(
                    tools["apply_data_changes"],
                    context,
                    {
                        "file": name,
                        "operation": "update",
                        "expected_revision": result["revision"],
                        "filters": [{"field": "code", "value": "A"}],
                        "values": [{"field": "amount", "value": 20}],
                    },
                )
                self.assertTrue(changed["success"], changed)
                audit = db.execute(
                    "SELECT key_value FROM project_data_changes ORDER BY id DESC LIMIT 1"
                ).fetchone()
                import json

                self.assertEqual(json.loads(audit["key_value"]), {"code": "A"})
                exported = await _invoke(
                    tools["export_dataset"], context, {"file": name}
                )
                self.assertTrue(exported["success"], exported)
                current = await _invoke(tools["query_dataset"], context, {"file": name})
                self.assertEqual(current["rows"][0]["amount"], 20)
                inserted = await _invoke(
                    tools["apply_data_changes"],
                    context,
                    {
                        "file": name,
                        "operation": "insert",
                        "expected_revision": current["revision"],
                        "rows": [
                            {
                                "values": [
                                    {"field": "code", "value": "B"},
                                    {"field": "amount", "value": 30},
                                ]
                            }
                        ],
                    },
                )
                self.assertTrue(inserted["success"], inserted)
                deleted = await _invoke(
                    tools["apply_data_changes"],
                    context,
                    {
                        "file": name,
                        "operation": "delete",
                        "expected_revision": inserted["revision"],
                        "filters": [{"field": "code", "value": "A"}],
                    },
                )
                self.assertTrue(deleted["success"], deleted)
                precise = await _invoke(
                    tools["export_dataset"], context, {"file": name}
                )
                self.assertFalse(precise["success"])
                template = await _invoke(
                    tools["export_dataset"], context, {"file": name, "mode": "template"}
                )
                self.assertTrue(template["success"], template)
                repeated = await _invoke(
                    tools["export_dataset"], context, {"file": name}
                )
                self.assertTrue(repeated["success"], repeated)
                self.assertEqual(repeated["written_cells"], 0)

            asyncio.run(flow())
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.active["A2"].value, "B")
            self.assertEqual(book.active["B2"].value, 30)
            book.close()
            # Editing the managed source makes the old import unusable.
            book = openpyxl.load_workbook(path)
            book.active["B2"] = 999
            book.save(path)
            book.close()
            stale = asyncio.run(
                _invoke(tools["query_dataset"], context, {"file": name})
            )
            self.assertFalse(stale["success"])
            self.assertIn("版本失效", stale["message"])
        db.close()


if __name__ == "__main__":
    unittest.main()
