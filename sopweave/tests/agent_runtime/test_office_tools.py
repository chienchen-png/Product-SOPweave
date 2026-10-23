from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document


class OfficeToolTests(unittest.TestCase):
    def test_text_document_cursor_reconstructs_unicode_without_loss(self):
        from tools.office_documents import read_document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'长文本.md'
            original = '财务✅\n'*13001
            path.write_text(original,encoding='utf8')
            cursor, parts = 0, []
            while True:
                page = read_document(path,character_offset=cursor)
                parts.append(page['text'])
                self.assertLessEqual(len(page['text']),20000)
                if not page['truncated']:
                    break
                self.assertGreater(page['next_character_offset'],cursor)
                cursor = page['next_character_offset']
            self.assertEqual(''.join(parts),original)

    def test_docx_pages_reconstruct_long_paragraph_table_and_headers(self):
        from tools.office_documents import read_document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "分页.docx"
            document = Document()
            long_text = "财务明细" * 12000
            document.add_paragraph(long_text)
            document.add_table(rows=1, cols=1).cell(0, 0).text = "表格内容"
            document.sections[0].header.paragraphs[0].text = "页眉"
            document.sections[0].footer.paragraphs[0].text = "页脚"
            document.save(path)
            first = read_document(path, limit=2)
            self.assertEqual(first['next_offset'], 0, 'a partial paragraph must not be skipped')
            offset, character_offset = 0, 0
            texts, versions = {}, set()
            for _ in range(10):
                page = read_document(path, offset=offset, limit=2, character_offset=character_offset)
                self.assertLessEqual(sum(len(b['text']) for b in page['blocks']), 20000)
                versions.add(page['version'])
                for block in page['blocks']:
                    texts[block['id']] = texts.get(block['id'], '') + block['text']
                if not page['truncated']:
                    break
                offset, character_offset = page['next_offset'], page['next_character_offset']
            else:
                self.fail('document pagination did not terminate')
            self.assertEqual(list(texts.values()), [long_text, '表格内容', '页眉', '页脚'])
            self.assertEqual(len(versions), 1)

    def test_shared_writeback_detects_concurrent_change(self):
        from tools.excel_writeback import load_writer

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "original"
            book.save(path)
            book.close()
            writer = load_writer(str(path))
            writer.set_cell("Sheet", 1, 1, "ours")
            book = openpyxl.load_workbook(path)
            book.active["A1"] = "external"
            book.save(path)
            book.close()
            try:
                with self.assertRaises(ValueError):
                    writer.save()
            finally:
                writer.close()
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.active["A1"].value, "external")
            book.close()

    def test_excel_coordinates_new_sheet_and_version_preserve_original(self):
        from tools.office_documents import read_sheet_range, edit_workbook
        from agent_runtime.workspace import file_version

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Original"
            sheet.append(["amount", "amount"])
            sheet.append([10, 20])
            sheet["C2"] = "=A2+B2"
            sheet["A2"].number_format = "0.00"
            book.save(path)
            book.close()
            result = read_sheet_range(path, "Original", "A2:B2", False)
            self.assertEqual([cell["value"] for cell in result["cells"]], [10, 20])
            version = file_version(path)
            edit_workbook(
                path,
                version,
                [{"sheet": "Added", "cell": "A1", "value": "new", "kind": "value"}],
            )
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.sheetnames, ["Original", "Added"])
            self.assertEqual(book["Original"]["C2"].value, "=A2+B2")
            self.assertEqual(book["Original"]["A2"].number_format, "0.00")
            book.close()
            with self.assertRaises(ValueError):
                edit_workbook(path, version, [])

    def test_docx_cross_run_replacement_keeps_format_and_table(self):
        from tools.office_documents import read_document, edit_document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.docx"
            document = Document()
            p = document.add_paragraph()
            p.add_run("收入").bold = True
            p.add_run("金额为10元").italic = True
            document.add_table(rows=1, cols=1).cell(0, 0).text = "keep"
            document.save(path)
            result = read_document(path)
            block = next(b for b in result["blocks"] if b["text"] == "收入金额为10元")
            edit_document(
                path,
                result["version"],
                [
                    {
                        "block_id": block["id"],
                        "old_text": "收入金额",
                        "new_text": "营业收入",
                    }
                ],
            )
            saved = Document(path)
            self.assertEqual(saved.paragraphs[0].text, "营业收入为10元")
            self.assertTrue(saved.paragraphs[0].runs[0].bold)
            self.assertTrue(saved.paragraphs[0].runs[-1].italic)
            self.assertEqual(saved.tables[0].cell(0, 0).text, "keep")
            changed_bytes = path.read_bytes()
            with self.assertRaises(ValueError):
                edit_document(path, result['version'], [{
                    'block_id':block['id'], 'old_text':'营业收入', 'new_text':'不应保存',
                }])
            self.assertEqual(path.read_bytes(), changed_bytes)


if __name__ == "__main__":
    unittest.main()
