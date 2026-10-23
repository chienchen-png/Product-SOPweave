from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from flask import Flask
from routes.project_data import project_data_bp

from storage.project_data_store import ProjectDataStore
from tools.project_data_importer import ProjectDataImporter, sha256_file
from tools.project_data_tools import export_project_data


class ExportWorkflowTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.path = self.root / "data.xlsx"
        book = openpyxl.Workbook()
        book.active.append(["Code", "Amount"])
        book.active.append(["001", 10])
        book.save(self.path)
        book.close()
        self.db = sqlite3.connect(":memory:")
        self.addCleanup(self.db.close)
        self.db.row_factory = sqlite3.Row
        self.db.executescript(Path("storage/schema.sql").read_text(encoding="utf-8"))
        result = ProjectDataImporter(self.db).import_workbook(
            "p", self.path.name, str(self.root), anchor_rows={"Sheet": 0}
        )
        self.file_id = result["file_id"]
        self.table = result["sheets"][0]["table_name"]
        self.store = ProjectDataStore(self.db)
        patcher = patch(
            "tools.project_data_tools._store", return_value=(self.store, self.db)
        )
        patcher.start()
        self.addCleanup(patcher.stop)

    def stage(self):
        self.db.execute(f'UPDATE "{self.table}" SET Amount=20 WHERE id=1')
        self.store.log_change(
            self.file_id,
            "Sheet",
            1,
            "update",
            col_index=1,
            col_name="Amount",
            old_value=10,
            new_value=20,
            key_value="001",
        )
        self.db.commit()

    def export(self, mode="precise"):
        return export_project_data(
            self.path.name, mode, project_id="p", work_dir=str(self.root)
        )

    def test_success_records_sync_checkpoint_without_deleting_audit_history(self):
        self.stage()
        first = self.export()
        self.assertTrue(first["success"], first)
        self.assertEqual(first["written_cells"], 1)
        self.assertEqual(self.store.count_changes(self.file_id)["total"], 1)
        self.assertEqual(
            self.store.count_changes(self.file_id, pending_only=True)["total"], 0
        )
        before = self.path.read_bytes()
        with patch("tools.project_data_tools.load_writer") as writer:
            repeated = self.export()
        self.assertTrue(repeated["success"], repeated)
        self.assertEqual(repeated["written_cells"], 0)
        writer.assert_not_called()
        self.assertEqual(self.path.read_bytes(), before)

    def test_changed_source_is_rejected_by_the_shared_export_entry(self):
        self.stage()
        book = openpyxl.load_workbook(self.path)
        book.active["B2"] = 999
        book.save(self.path)
        book.close()
        before = self.path.read_bytes()
        result = self.export()
        self.assertFalse(result["success"], result)
        self.assertRegex(result["error"], "更改|变化|版本")
        self.assertEqual(self.path.read_bytes(), before)

    def test_precise_export_refuses_deletions_instead_of_skipping_them(self):
        self.store.log_change(
            self.file_id, "Sheet", 1, "delete", col_index=0, old_value="001"
        )
        self.db.commit()
        before = self.path.read_bytes()
        result = self.export()
        self.assertFalse(result["success"], result)
        self.assertRegex(result["error"], "template|重建")
        self.assertEqual(self.path.read_bytes(), before)

    def test_precise_export_refuses_more_than_2000_pending_changes(self):
        for index in range(2001):
            self.store.log_change(
                self.file_id,
                "Sheet",
                1,
                "update",
                col_index=1,
                old_value=10,
                new_value=20,
            )
        self.db.commit()
        result = self.export()
        self.assertFalse(result["success"], result)
        self.assertIn("2000", result["error"])

    def test_failed_save_does_not_mark_changes_synced(self):
        self.stage()
        before = sha256_file(self.path)
        with patch("tools.project_data_tools.load_writer") as factory:
            factory.return_value.save.side_effect = OSError("injected save failure")
            result = self.export()
        self.assertFalse(result["success"], result)
        meta = json.loads(self.store.get_file(self.file_id)["meta_json"])
        self.assertEqual(meta.get("exported_change_id", 0), 0)
        self.assertEqual(sha256_file(self.path), before)

    def test_ui_counters_separate_pending_changes_from_retained_history(self):
        self.stage()
        self.assertTrue(self.export()["success"])
        app = Flask(__name__)
        app.register_blueprint(project_data_bp)
        with patch("routes.project_data.get_db", return_value=self.db):
            listing = (
                app.test_client()
                .get("/api/project-data/list?project_id=p")
                .get_json()["data"]["files"][0]
            )
            history = (
                app.test_client()
                .get(f"/api/project-data/{self.file_id}/changes")
                .get_json()["data"]
            )
        self.assertEqual(listing["changes_count"]["total"], 0)
        self.assertEqual(history["summary"]["total"], 1)
        self.assertEqual(history["pending_summary"]["total"], 0)
        self.assertTrue(history["changes"][0]["applied"])

    def test_insert_export_then_delete_rebuild_keeps_physical_lineage_without_ghost_rows(
        self,
    ):
        for code, amount in [("002", 30), ("003", 40)]:
            inserted = self.db.execute(
                f'INSERT INTO "{self.table}" (Code,Amount) VALUES (?,?)', (code, amount)
            ).lastrowid
            for column, value in enumerate((code, amount)):
                self.store.log_change(
                    self.file_id,
                    "Sheet",
                    inserted,
                    "insert",
                    col_index=column,
                    new_value=value,
                )
            self.db.commit()
            self.assertTrue(self.export()["success"])
        self.db.execute(f'DELETE FROM "{self.table}" WHERE id=2')
        self.store.log_change(
            self.file_id, "Sheet", 2, "delete", col_index=0, old_value="002"
        )
        self.db.commit()
        result = self.export("template")
        self.assertTrue(result["success"], result)
        book = openpyxl.load_workbook(self.path)
        try:
            values = [
                list(row)
                for row in book.active.iter_rows(min_row=2, max_row=4, values_only=True)
            ]
        finally:
            book.close()
        self.assertEqual(values, [["001", 10], [None, None], ["003", 40]])
        self.assertEqual(
            [
                tuple(row)
                for row in self.db.execute(
                    f'SELECT Code,_source_row FROM "{self.table}" ORDER BY id'
                )
            ],
            [("001", 2), ("003", 4)],
        )
